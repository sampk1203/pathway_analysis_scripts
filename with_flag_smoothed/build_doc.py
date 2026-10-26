#!/usr/bin/env python3
"""build_doc.py - one docx from many dopant folders, processed in groups.

python build_doc.py {A B C} {D E Ga=Li24d} {F G H Al=Li24d Ca=La} -o report.docx

Each {...} = one group. Per folder: heading + images. After all folders of a group: comparison tables
(same tables summ.py --compare writes).

Site deviation (optional, per group): add  <dopant element>=<sublattice>  tokens at the END of a group, e.g.
    {1Al-doped 2Al-doped Al=Li24d}      sublattices: Li24d  Li96h  Li  La  Zr  O
Element of LAMMPS atom type 1..N (if the folders have no element_list file), per group, also at the end:
    {1Al-doped 2Al-doped Al=Li24d elements=Li,La,Zr,Al,O}      (a global --elements Li La Zr Al O also works)
The dopant is matched to the closest site of that sublattice of the pure Ia-3d LLZO CIF; La, Zr, O are matched one-to-one
(closest free site) on their time-mean positions, Li per frame to the nearest free Li site. Shells = 05's (distance to the
dopant; Li: dopant position per frame, smoothed over --dop-smooth-ps; La/Zr/O: time-mean). Needs the raw dumps + element_list in each folder. Adds, per folder: one page with a violin figure and one
projection figure per dopant atom for every T (after the msd05 image), and, per group, deviation tables after the comparison.
Figures, csv and a cache (npz) go to <folder>/sitedev/<folder name>/ ; --recompute ignores the cache.

RMSD (every folder, no dopant spec needed): RMSD of every atom type vs the first frame of the window (framework drift removed).
Grid: one row per atom type, one column per T; thin line = single atom, black = species RMSD; x = frame, y = RMSD (A), y shared
by every panel (auto-fitted). Inserted right after the first image of each folder. Needs the raw dumps in
<folder>/MD-<T>/ (as the site deviation). Figure + cache go to <folder>/rmsd/<folder name>/ ; --recompute ignores the cache.

Local Li density bins (every folder with a dopant, no spec needed): does Li MSD / correlation depend on the LOCAL Li density at fixed
distance from the dopant? Local density of a Li at t0 = number of other Li within --rho-r A (default 4), also given relative to the
cell-mean value (1 = average). The bins are set ONCE per T from the far Li (farther than the last shell edge from EVERY dopant at t0, no
dopant proximity in the reference): --rho-bins bins of fixed width in relative density over 0.1-1.5 (RHO_LO, RHO_HI; first/last bin open-ended; not equal-count, sparse bins may be empty). The SAME
bin edges are then applied to the Li of every shell (0-3, 3-5, ... , far), so shells are compared at equal local density. Per shell, bin and
lag tau (--rho-lags, ps, default 1 5 10):
    MSD(tau)  = < |r(t0+tau) - r(t0)|^2 >                                   over the Li of the bin (shell rule 'start', as 05)
    H_R,loc   = sum_i sum_{j in S_i} |dr_j|^2  /  sum_i |sum_{j in S_i} dr_j|^2,   S_i = Li i + its Li neighbours within --rho-r at t0
H_R,loc = 1: neighbouring Li move independently; < 1: their displacements are positively correlated (concerted). It is NOT 09's
intra-shell H_R: Li of one density bin are spread over the cell, so their summed displacement is uncorrelated by construction (-> 1);
the neighbourhood sum measures correlation where the density is. Overlapping neighbourhoods share Li, so the SE (delete-one-block
jackknife, 5 time blocks) is only a rough scale. msd_top_over_bottom (csv) = MSD of the densest bin / MSD of the sparsest bin:
~1 -> density does not set MSD at fixed distance (proximity does); > 1 -> denser = faster (densest / sparsest bin with >= 30 samples).
Figures, added after the msd05 image: rho_bins.png = far Li, row 1 MSD, row 2 H_R,loc, x = relative local density, one line per T, one
column per tau; rho_shells_T<T>K.png (one per T) = same rows/columns, one line per shell; open symbol = fewer than 3 Li of that bin per
origin (noisy, H_R -> 1 by construction). Near shells hold few Li: bins that no Li of the shell falls in are empty (no point). Figures, csv (rho_bins.csv) and cache go to
<folder>/rhobin/<folder name>/ ; --recompute ignores the cache. No sites, no hop definition; needs the raw dumps (as RMSD).

Shell density vs MSD / H_R, pooled over all folders (every folder with a dopant, no spec needed): is the faster / more concerted shell
faster because it holds more Li, or because of where it is (distance to the dopant, strain, site mix)? Per folder, T, shell (as 05/09:
shell = distance to the nearest dopant at t0, 'start' ownership) and lag tau (--pool-lags, ps): Li density in the shell relative to a uniform
density (05's density ratio), MSD, MSD / MSD(far) (05), H_R of the Li of the shell (09's intra-shell Haven ratio). Figure per T (pool_T<T>K.png,
next to the report in shellpool/): rows = MSD, MSD/far, H_R; columns = tau; x = shell density ratio; one dot per (folder, shell), colour = shell,
marker = folder, pale dot = fewer than --pool-min-n Li in the shell (H_R -> 1 by construction). The number printed in each panel is the Spearman
rho between density and the y value ACROSS FOLDERS, WITHIN one shell (n = folders; needs >= 4): if it is clearly non-zero, density predicts
MSD / H_R at fixed distance to the dopant; if only the colour-to-colour (between-shell) trend exists, density is confounded with distance.
'all shells' = rho over every dot (between + within). shell_corr.csv has the numbers, shell_pool.csv every point.
CONTROL (diamonds): the same shells around --ctrl-n reference points that are NOT dopants: positions of random Li atoms at frame 0
(--ctrl-ref li, a Li-site-like point; --ctrl-ref random = random points), each at least the last shell edge away from every real dopant; Li
within the last shell edge of a real dopant are left out of the control (so no real dopant sits inside a control shell). Diamond = mean over the
control points. If the control shows the same shell pattern, the pattern is lattice geometry, not the dopant. A random
Li-site point is not identical to a dopant on a site (no charge compensation, no strain): a null for geometry only. Cache: <folder>/shellpool/.
Needs the raw dumps (as RMSD); 05_msd_by_shell.py and 09_shell_haven.py from --scripts are imported, so the numbers are theirs.

Radius of influence per dopant element (every folder with a dopant, no spec needed; NOT in the report yet, files only): no shells. Every Li at
every t0 gives one point (r = distance to the NEAREST dopant at t0, MSD over tau, 'start' rule as 05). Gaussian-kernel regression
(--rad-bw, default 0.5 A, grid --rad-step over 0..--rad-rmax) gives the profile y(r) = MSD(r) / MSD(far), far = Li farther than --rad-far
(default: last shell edge) from every dopant. Folders of one dopant element (1Al, 2Al, 3Al) are pooled; 95 % CI-free: the band is a 16-84 %
block bootstrap (time blocks resampled per folder, --rad-boot). Radius = three fits with a FREE plateau B (B > 1: MSD still rising past the far edge,
the radius is larger than the far edge and the number is a lower bound): exp B(1 - A e^(-r/lam)), gauss B(1 - A e^(-(r/lam)^2)),
sigmoid B(1 - A / (1 + e^((r - r0)/w))); reported: r95 / r90 (deficit down to 5 / 10 % of the plateau) and r50 (half recovery), best fit = lowest
AIC (kernel points overlap, so AIC is only a rough guide). Model-free: r_noise = last radius where y is below 1 - 2 sd of the control profiles;
ell = integral of the deficit (1 - y) dr; V = integral of 4 pi r^2 (1 - y) dr; r_eff = radius of the fully frozen sphere with the same V.
Control = the same profile around --ctrl-n non-dopant Li-site points (only Li far from every real dopant), so it is 1 on average; its scatter is
the null band. Caps: the radius cannot be measured beyond half the box or the dopant spacing (radius_summary.csv: half_box_A, dopant_dopant_min_A,
r_p99_A = 99th percentile of the nearest-dopant distance). Files in <report dir>/radius/: radius_summary.csv (every fit, element and folder
level), radius_profiles.csv, radius_profiles_T<T>K.png (rows tau, columns element), radius_summary_tau<tau>ps.png. Cache: <folder>/radius/.
--radius-only runs just this (nothing else is recomputed, no docx / tex / pdf); --no-radius skips it.
Dopant-comparison figures (both radius/ and radius_plainmsd/): radius_allT_tau<tau>ps.png = r95, depth A and plateau B of the sigmoid fit
for ALL temperatures at one tau (x = dopant element, colour = T); radius_allT_alltau.png = all temperatures AND all taus (marker = tau). Only
valid fits are drawn: r95 inside the measured range (no '*') and A >= 0.1. The temperature-dependence panel of radius_summary shows the same valid
sigmoid r95 only, one line per element.
PLAIN-MSD FIT (same kernel sums, so no extra trajectory pass): the identical analysis, but the profile fitted is the plain MSD(r) in A^2, NOT
divided by MSD(far). Written to <report dir>/radius_plainmsd/ (same file names as radius/, so the two can be compared side by side; profile csv has
msd_A2 instead of ratio, plus far_msd_A2 for reference). The plateau B is then in A^2 and free (0.5 - 3 x the mean of the outermost 3 fitted points);
A, r50, r90, r95 are relative to the plateau, so they compare directly with the ratio fit. Null band = mean +- 2 sd of the control profiles (A^2).
Model-free r_noise / ell / V / r_eff need a dimensionless deficit: they use y / (mean of the outermost 5 valid points), i.e. the profile's own plateau
instead of the far-Li MSD. --no-plain skips this fit.
Fits use only r >= --rad-rmin (default 2.5 A): the profile below that has few Li and a spike (selection artefact), so it is not fitted (still drawn;
r_eff / r_noise use the whole profile). The summary figure shows the sigmoid fit only (exp / gauss stay in the csv and the profile figure).
radius_vs_conc_T<T>K_tau<tau>ps.png: per-FOLDER sigmoid r95 and A vs number of dopant atoms, one line per element (folder level, no bootstrap;
r95 is not drawn where A < 0.1, there it is noise).

LaTeX (always): next to the report, <out>.tex + figures1/ with every figure of the report (svg; pdf made from it with rsvg-convert
or inkscape; png only where no vector version exists). Same order as the docx, with the same tables (comparison + deviation). Add your text around
the figures in the .tex and compile it yourself. Figures are included without file extension, so pdflatex takes the .pdf.

Hop measures (computed here, from hops_cif/<name>/hops_T<T>K.csv + the dumps, cached in <folder>/commit/): %returned, committed hops,
COMMIT/far (--tau-ps, default 1 ps; frames = tau / dump frame spacing). Table columns of 07's loop erasure are tagged 'static dopant calc.'.

Needs summ.py next to this file, the analysis scripts (llzo_io.py, llzo_sites.py ...) in ./with_flag (or --scripts DIR),
python-docx, numpy, scipy, matplotlib. Quote braces if the shell eats them:  python build_doc.py '{A B}' '{C D}'
Put options after the groups.
"""
import argparse, csv, dataclasses, glob, json, os, re, shutil, subprocess, sys

HERE = os.path.dirname(os.path.abspath(__file__))


def _scripts_dir():
    a = sys.argv
    for i, x in enumerate(a):
        if x == '--scripts' and i + 1 < len(a): return a[i + 1]
        if x.startswith('--scripts='): return x.split('=', 1)[1]
    return os.path.join(HERE, 'with_flag')


SCRIPTS = os.path.abspath(_scripts_dir())
sys.path.insert(0, HERE)
sys.path.insert(0, SCRIPTS)
import numpy as np
import summ
import llzo_io as lio
import llzo_sites as ls
import matplotlib.pyplot as plt
from scipy.optimize import linear_sum_assignment
from scipy.spatial import cKDTree
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PAGE_W, PAGE_H, MARGIN = 29.7, 21.0, 1.27          # A4 landscape, cm
USABLE_W = PAGE_W - 2 * MARGIN
MAX_H = PAGE_H - 2 * MARGIN - 1.0                  # keep images inside one page
CIF_DEFAULT = '/media/qpivolta/8TB/archive_vm/backup-26-3-26-voltaet/2_debdipto/2x2x4-new/pathway_analysis_scripts/cubic_LLZO_wyckoff.cif'

# (layout, subfolder, filename pattern); 'row' = all matches in one line, 'each' = one per line, page width
STEPS = [
    ('row',  'results',  'diff_T_kT_*.png'),
    ('each', 'results',  'vanhove_T800K.png'),
    ('each', 'haven09',  'haven_shell.png'),
    ('each', 'hops_cif', 'hop_network_T800K.png'),
    ('each', 'hops_cif', 'hop_network_T800K_*.png'),
    ('each', 'hops_cif', 'hops_summary.png'),
    ('each', 'hops_cif', 'concerted_vs_chance_shell.png'),
    ('each', 'hops_cif', 'refill_vs_chance_shell.png'),
    ('each', 'loops07',  'hop_loops.png'),
    ('each', 'msd05',    'msd_shell_ratio.png'),
]


SVG_MODE = False                 # False: PNG report (docx). True: SVG report (-> pdf)
SVG_MISSING, _SVG_N = [], [0]


def warn(msg):
    if not SVG_MODE: print('warning: ' + msg, file=sys.stderr)       # 2nd pass repeats the same warnings: stay quiet


_SVG_EXT = ('<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:ext uri="{{96DAC541-7B7A-43D3-8B79-37D633B846F1}}">'
            '<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="{rid}"/></a:ext></a:extLst>')


def add_pic(run, path, width):
    """PNG picture; in SVG_MODE also attach the sibling .svg (PNG stays as fallback, so Word and LibreOffice both work)."""
    shape = run.add_picture(path, width=width)
    if SVG_MODE:
        sv = os.path.splitext(path)[0] + '.svg'
        if os.path.exists(sv):
            _SVG_N[0] += 1
            with open(sv, 'rb') as fh:
                part = Part(PackURI(f'/word/media/svg{_SVG_N[0]}.svg'), 'image/svg+xml', fh.read(), run.part.package)
            rid = run.part.relate_to(part, RT.IMAGE)
            shape._inline.xpath('.//a:blip')[0].append(parse_xml(_SVG_EXT.format(rid=rid)))
        else:
            SVG_MISSING.append(path)
    return shape
def nat(s): return [int(t) if t.isdigit() else t.lower() for t in re.split(r'(\d+)', s)]
def fname(folder): return os.path.basename(os.path.normpath(folder))


# ---------------------------------------------------------------- args
SPEC_RX = re.compile(r'^[A-Z][a-z]?=[A-Za-z0-9_-]+$')
ELEM_RX = re.compile(r'^elements=([A-Za-z,]+)$')       # elements=Li,La,Zr,Al,O : element of LAMMPS type 1..N for the whole group
SUBLAT = {'li24d': ('Li', '24d'), 'li96h': ('Li', '96h'), 'li': ('Li', None), 'la': ('La', None),
          'zr': ('Zr', None), 'o': ('O', None)}


def sublattice(txt):
    k = txt.lower().replace('-', '').replace('_', '')
    if k not in SUBLAT: sys.exit(f"unknown sublattice '{txt}' (use Li24d, Li96h, Li, La, Zr, O)")
    return SUBLAT[k]


def parse_groups(items):
    """-> [{'folders': [...], 'sites': {'Al': 'Li24d'}}]; tokens like Al=Li24d are dopant-site specs, not folders."""
    groups, cur = [], None

    def add(w):
        m = ELEM_RX.match(w)
        if m:
            cur['elements'] = [x.capitalize() for x in m.group(1).split(',') if x]
        elif SPEC_RX.match(w):
            el, sub = w.split('=')
            sublattice(sub)
            cur['sites'][el] = sub
        else:
            cur['folders'].append(w)

    for tok in items:
        for part in re.split(r'([{}])', tok):
            if part == '{':
                if cur is not None: sys.exit('nested { not allowed')
                cur = {'folders': [], 'sites': {}, 'elements': None}
            elif part == '}':
                if cur is None: sys.exit('unmatched }')
                groups.append(cur); cur = None
            elif part.strip():
                if cur is None: sys.exit(f'folder outside braces: {part}')
                # '{A B}' quoted as one arg -> split on spaces; a plain quoted "my dir" stays whole
                for w in (part.split() if ('{' in tok or '}' in tok) else [part.strip()]): add(w)
    if cur is not None: sys.exit('missing closing }')
    return [g for g in groups if g['folders']]


def expand(group):
    """shell does not expand globs inside {...}; do it here (folders only, natural order)"""
    out = []
    for f in group:
        if any(c in f for c in '*?'):
            hit = sorted((p for p in glob.glob(f) if os.path.isdir(p)), key=lambda p: nat(fname(p)))
            if not hit: warn(f'pattern matched no folder: {f}')
            out += hit
        else:
            out.append(f)
    return out


# ---------------------------------------------------------------- images
def find(folder, sub, pat):
    base = glob.escape(os.path.join(folder, sub))
    return sorted(glob.glob(os.path.join(base, '**', pat), recursive=True),
                  key=lambda p: nat(os.path.basename(p)))


def aspect(path):
    im = Image.from_file(path)
    return im.px_width / im.px_height


def fit(path, w, h):
    im = Image.from_file(path)
    s = min(w / im.px_width, h / im.px_height)
    return im.px_width * s


def add_images(doc, paths, row, max_h=MAX_H):
    if row:
        # borderless 1-row table: fixed cell widths, so images can never wrap onto a 2nd line
        cw = USABLE_W / len(paths)
        t = doc.add_table(rows=1, cols=len(paths)); t.autofit = False
        for cell, path in zip(t.rows[0].cells, paths):
            cell.width = Cm(cw)
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            add_pic(p.add_run(), path, Cm(fit(path, cw - 0.5, max_h)))
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    else:
        for path in paths:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            add_pic(p.add_run(), path, Cm(fit(path, USABLE_W, max_h)))


def add_sitedev_page(doc, name, rows, title='Site deviation from pure LLZO'):
    """rows: one list of image paths per temperature, same length. ALL on one page: one table, every image the same height."""
    h = doc.add_heading(f'{title}: {name}', 2)
    h.paragraph_format.page_break_before = True
    h.paragraph_format.space_after = Pt(2)
    nrow, ncol = len(rows), len(rows[0])
    asp = [[aspect(p) for p in r] for r in rows]
    ca = [max(a[i] for a in asp) for i in range(ncol)]
    hh = min((USABLE_W - 0.6 * ncol) / sum(ca), (PAGE_H - 2 * MARGIN - 2.0) / nrow - 0.15)
    t = doc.add_table(rows=nrow, cols=ncol); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for c in range(ncol): t.columns[c].width = Cm(ca[c] * hh + 0.6)       # tblGrid too: LibreOffice uses it, else it clips the wide violin
    for r, (paths, a) in enumerate(zip(rows, asp)):
        for c in range(ncol):
            cell = t.cell(r, c); cell.width = Cm(ca[c] * hh + 0.6)
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            add_pic(p.add_run(), paths[c], Cm(a[c] * hh))


def add_folder(doc, folder, page_break, sd=None, rmsd=None, rho=None):
    h = doc.add_heading(fname(folder), 1)
    h.paragraph_format.page_break_before = page_break
    first_img = True
    for layout, sub, pat in STEPS:
        paths = find(folder, sub, pat)
        if not paths:
            miss = os.path.join(folder, sub, '**', pat)
            warn(f'missing: {miss}')
            doc.add_paragraph(f'[missing: {miss}]')
            continue
        # first image block shares its page with the folder heading -> smaller, else it is pushed to the next page
        add_images(doc, paths, layout == 'row', MAX_H - 4.0 if first_img else MAX_H)
        if first_img and rmsd: add_images(doc, [rmsd], False, MAX_H)            # RMSD grid right after the very first image
        first_img = False
    if rho: add_images(doc, rho, False, MAX_H)                                  # local-density bins, after the msd05 image
    if sd:
        add_sitedev_page(doc, fname(folder), [[r['violin']] + r['projs'] for r in sd])
        add_sitedev_page(doc, fname(folder), [r['full'] for r in sd], title='Whole-cell projection')


# ---------------------------------------------------------------- tables
def shade(cell, color):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shd)


def put(cell, text, pt, head=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.font.name = 'Arial'; r.font.size = Pt(pt); r.bold = head
    if head: shade(cell, 'E8E8E8')


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)


STATIC = '[static dopant calc.] '
FEW = 20


def grid_data(defn, cols_by_sys, Ts, rows, rowname, getter):
    """Everything a table shows, renderer-free: docx and tex are both drawn from this, so the numbers are identical."""
    cols = [(s, a) for s, _, atoms in cols_by_sys for a in atoms]
    return dict(defn=defn, syshead=[(lab, len(at)) for _, lab, at in cols_by_sys if at],
                atoms=[summ.dop_name(a) for _, a in cols], Ts=list(Ts), rowlabels=[rowname(r) for r in rows],
                cells=[[[getter(s, T, a, r) for s, a in cols] for r in rows] for T in Ts])


def render_grid(doc, d):
    nd, ns, Ts = len(d['atoms']), len(d['rowlabels']), d['Ts']
    nc = 2 + nd
    pt = 8 if nc <= 10 else 7 if nc <= 16 else 6
    t = doc.add_table(rows=3 + len(Ts) * ns, cols=nc)
    t.style = 'Table Grid'; t.autofit = False
    wid = [1.7, 1.9] + [(USABLE_W - 3.6) / nd] * nd
    for row in t.rows:
        for i, c in enumerate(row.cells): c.width = Cm(wid[i])
    put(t.cell(0, 0), 'Temperature (K)', pt, True); t.cell(0, 0).merge(t.cell(2, 0))
    put(t.cell(0, 1), 'Shell', pt, True); t.cell(0, 1).merge(t.cell(2, 1))
    put(t.cell(0, 2), d['defn'], pt, True)
    if nd > 1: t.cell(0, 2).merge(t.cell(0, nc - 1))
    c = 2
    for label, n in d['syshead']:
        put(t.cell(1, c), label, pt, True)
        if n > 1: t.cell(1, c).merge(t.cell(1, c + n - 1))
        c += n
    for j, a in enumerate(d['atoms']): put(t.cell(2, 2 + j), a, pt, True)
    for r in range(3): repeat_header(t.rows[r])
    r = 3
    for ti, T in enumerate(Ts):
        for i, lab in enumerate(d['rowlabels']):
            if i == 0:
                put(t.cell(r, 0), T, pt)
                if ns > 1: t.cell(r, 0).merge(t.cell(r + ns - 1, 0))
            put(t.cell(r, 1), lab, pt)
            for j, txt in enumerate(d['cells'][ti][i]): put(t.cell(r, 2 + j), txt, pt)
            r += 1
    doc.add_paragraph()


def render_docx(doc, blocks):
    """blocks: ('h', level, text, page_break) | ('p', text) | ('t', grid_data)"""
    for b in blocks:
        if b[0] == 'h':
            h = doc.add_heading(b[2], b[1])
            if b[3]: h.paragraph_format.page_break_before = True
        elif b[0] == 'p': doc.add_paragraph(b[1])
        else: render_grid(doc, b[1])


def cmp_table(defn, systems, src, getter):
    Ts = sorted(set().union(*[set(s['d' + src]) for s in systems]))
    return ('t', grid_data(defn, [(s, s['name'], summ.atoms_of(s, src)) for s in systems], Ts,
                           summ.shell_union(systems, src), summ.shell_name, getter))


def hop_defs(systems):
    """[(key, definition)]: reliable columns first, then 07's loop-erasure columns tagged static dopant calc."""
    D = dict(summ.defs07(systems[0]['meta']))
    cm = [s['commit'] for s in systems if s.get('commit')]
    tau = ' / '.join(sorted({f"{c['tau_ps']:g} ps ({c['tau_frames']} frames of {c['dt']:.3f} ps)" for c in cm})) or 'tau'
    new = [('hops', D['hops']), ('allfar', D['allfar']),
           ('ret', f'% of hops undone: the same Li hops back into the origin site of the hop within tau = {tau} (± SE)'),
           ('ncommit', f'Committed hops: hops not undone within tau = {tau}'),
           ('commitfar', 'Committed-hop rate per Li in shell relative to far shell (>7 Å), ± SE (1 = same as far; <1 = fewer lasting hops)')]
    old = [(k, STATIC + D[k]) for k in ('loop', 'palin', 'nonpal', 'net', 'robust', 'netfar', 'robfar', 'smallnet', 'smallloop')]
    return new + old


def hop_notes(systems):
    return [
        'Reliable hop measures (columns without a tag). Each Li is followed along its WHOLE hop path; nothing is cut at shell edges. '
        'A hop is assigned to the shell of the Li (distance to the dopant at the hop frame, as all/far). '
        'Returned: the same Li makes a later hop, at most tau after it, that ends in the origin site of this hop (a→b, then …→a). '
        'Committed = not returned. tau default 1 ps, converted to frames with the frame spacing of each dump. '
        '%returned = 100 × returned / hops. COMMIT/far = (committed hops / Li-time in shell) / (same for far Li); '
        'all/far = same with all hops. Reading: all/far high with %returned high = flicker; COMMIT/far is the lasting-hop rate.',
        '± = delete-one-block jackknife SE over 5 time blocks of one run; * = fewer than 20 events; n/a = undefined. '
        'Hops within tau of the end of the run cannot be seen returning (counted committed in every shell alike).',
        'Columns tagged [static dopant calc.] are 07 as run: loop erasure with the path cut at every shell change, shell from the '
        'per-frame (smoothed) dopant position. Kept for continuity. Their near-shell NET/far and ROBUST/far trends changed with the '
        'dopant smoothing window (0 vs 1 ps): do not read them as transport; use COMMIT/far.']


def compare_blocks(systems):
    B = [('h', 1, 'Comparison', True),
         ('p', 'Each table: rows = temperature and shell; columns = dopant system, then dopant atom '
               '(average = nearest dopant to each Li). "–" = not available.')]

    def g05(s, T, a): return s['d05'].get(T, {}).get(a, {})

    taus = sorted({t for s in systems for T in s['d05'].values() for v in T.values() for t in v['msd']}, key=float)
    B += [('h', 2, 'Li density around dopant', False),
          cmp_table(summ.DENS_DEF, systems, '05', lambda s, T, a, sh: summ.dens(g05(s, T, a).get('shells', {}).get(sh)))]
    for t in taus:
        B += [('h', 2, f'Li MSD, lag τ = {t} ps', False),
              cmp_table(summ.MSD_DEF + f' (τ = {t} ps)', systems, '05',
                        lambda s, T, a, sh, t=t: summ.msdc(g05(s, T, a).get('msd', {}).get(t, {}).get(sh))),
              cmp_table(summ.RAT_DEF + f' (τ = {t} ps)', systems, '05',
                        lambda s, T, a, sh, t=t: summ.ratc(g05(s, T, a).get('msd', {}).get(t, {}).get(sh)))]
    B.append(('h', 2, 'Hop loops around dopant', False))
    B += [('p', x) for x in hop_notes(systems)]
    for k, txt in hop_defs(systems):
        B.append(cmp_table(txt, systems, '07',
                           lambda s, T, a, sh, k=k: summ.fmt(s['d07'].get(T, {}).get(a, {}).get(sh, {}).get(k))))
    return B


# ---------------------------------------------------------------- reliable hop measures (returned / committed hops)
def returned_flags(L_, f_, sa, sb, tau):
    """hop k of a Li is 'returned' if a later hop of the SAME Li, <= tau frames after it, ends in site_from of hop k."""
    ret = np.zeros(len(L_), bool)
    order = np.lexsort((f_, L_))
    for g in np.split(order, np.flatnonzero(np.diff(L_[order])) + 1):
        for k in range(len(g)):
            for j in range(k + 1, len(g)):
                if f_[g[j]] - f_[g[k]] > tau: break
                if sb[g[j]] == sa[g[k]]:
                    ret[g[k]] = True; break
    return ret


def commit_rows(folder, cfg, T, p, tau_ps, hops_csv):
    traj = lio.prepare(cfg, md_dump(folder, T) or lio.dump_path(cfg, T), p['tmin'], verbose=False)
    drift = lio.framework_drift(cfg, traj)
    li = traj.pos[:, lio.li_mask(cfg, traj), :] - drift[:, None, :]
    F = li.shape[0]
    dt = float(np.median(np.diff(traj.time_ps)))
    tau = max(1, int(round(tau_ps / dt)))
    dops, cnt = {}, {}
    for t in cfg['dopant_types']:
        for i in np.where(traj.types == t)[0]:
            el = cfg['types'][t]
            cnt[el] = cnt.get(el, 0) + 1
            dops[f'{el}{cnt[el]}'] = lio.dopant_xyz(cfg, traj, drift, i)
    if not dops: raise ValueError('no dopant atoms')
    dists = {lab: lio.dist_to(li, x, traj.L) for lab, x in dops.items()}
    dnear = np.min(list(dists.values()), axis=0)
    dists = {'any': dnear, **dists}
    H = list(csv.DictReader(open(hops_csv)))
    if not H: return []
    L_ = np.array([int(h['li']) for h in H]); f_ = np.array([int(h['frame']) for h in H])
    sa = np.array([int(h['site_from']) for h in H]); sb = np.array([int(h['site_to']) for h in H])
    ret = returned_flags(L_, f_, sa, sb, tau)
    edges = np.array(p['edges'], float)
    names = lio.shell_names(edges)
    K, nb = len(names), lio.DEFAULT_BLOCKS
    rows = []
    for lab, d_ in dists.items():
        sh = lio.shell_labels(d_, dnear, edges, 'all')
        C = lio.block_counts(sh, K, nb)
        res = lio.hop_shell_rows(T, lab, 'returned', names, sh[f_, L_], f_, ret, np.ones(len(H), bool),
                                 litime_blk=C * dt, nmean=C.sum(0) / F, F=F, nblocks=nb)
        for k, x in enumerate(res):
            nc = x['n_hops'] - x['n_back']
            rows.append(dict(T=int(T), dopant=lab, shell=names[k], is_far=bool(k == K - 1), n_hops=int(x['n_hops']),
                             n_commit=int(nc), ret_pct=float(x['back_pct']), ret_se=float(x['back_pct_se']),
                             commit_ratio=float(x['nonback_ratio_to_far']),
                             commit_se=float(x['nonback_ratio_se']) if nc >= lio.MIN_N else float('nan'),
                             tau_ps=float(tau_ps), tau_frames=tau, dt=dt))
    return rows


def _pmstar(v, se, n, dec, ref=False):
    if ref: return '1 (ref)'
    if not np.isfinite(v): return 'n/a'
    return f'{v:.{dec}f} ± ' + (f'{se:.{dec}f}' if np.isfinite(se) else 'nan') + ('*' if n < FEW else '')


def merge_commit(system, rows):
    for r in rows:
        if not r['n_hops']: continue
        c = system['d07'].setdefault(r['T'], {}).setdefault(r['dopant'], {}).setdefault(r['shell'], {})
        c['ret'] = _pmstar(r['ret_pct'], r['ret_se'], r['n_hops'], 1)
        c['ncommit'] = str(r['n_commit'])
        c['commitfar'] = _pmstar(r['commit_ratio'], r['commit_se'], r['n_commit'], 2, r['is_far'])
        system['commit'] = dict(tau_ps=r['tau_ps'], tau_frames=r['tau_frames'], dt=r['dt'])


def _commit_task(args):
    f, cfg, T, p, tau_ps, hops, outp, key = args
    try:
        rows = commit_rows(f, cfg, T, p, tau_ps, hops)
        os.makedirs(os.path.dirname(outp), exist_ok=True)
        with open(outp, 'w') as fh: json.dump({'key': key, 'rows': rows}, fh)
    except Exception as e:
        return f'{f} T={T}: {e}'


def run_commit(groups, a, p):
    """-> {(group index, folder): [rows]} (cached in <folder>/commit/)."""
    todo, plan = [], {}
    for gi, g in enumerate(groups):
        for f in g['folders']:
            if not os.path.isdir(f): continue
            try:
                ga = argparse.Namespace(**vars(a))
                if g['elements']: ga.elements = g['elements']
                cfg = lio.load_config(f, None, ga)
                temps = a.T or md_temps(f) or lio.temperatures(cfg)
            except SystemExit as e:
                warn(f'commit measures skipped for {f}: {e}'); continue
            for T in temps:
                hops = os.path.join(f, 'hops_cif', cfg['name'], f'hops_T{T}K.csv')
                if not os.path.exists(hops):
                    alt = sorted(glob.glob(os.path.join(f, 'hops_cif', '*', f'hops_T{T}K.csv')))
                    if not alt: warn(f'no hops csv for {f} T={T}'); continue
                    hops = alt[0]
                dump = md_dump(f, T) or lio.dump_path(cfg, T)
                key = repr((a.tau_ps, tuple(p['edges']), p['tmin'], cfg.get('dop_smooth_ps'), tuple(sorted(cfg['types'].items())),
                            os.path.getmtime(hops), os.path.getmtime(dump), lio.DEFAULT_BLOCKS, 'commit-v1'))
                outp = os.path.join(f, 'commit', f'commit_T{T}K.json')
                plan.setdefault((gi, f), []).append(outp)
                ok = False
                if os.path.exists(outp) and not a.recompute:
                    try: ok = json.load(open(outp))['key'] == key
                    except Exception: ok = False
                if not ok: todo.append((f, cfg, T, p, a.tau_ps, hops, outp, key))
    print(f'[commit measures] {len(todo)} (folder, T) to compute, {sum(len(v) for v in plan.values()) - len(todo)} cached')
    if a.jobs > 1 and len(todo) > 1:
        from concurrent.futures import ProcessPoolExecutor
        with ProcessPoolExecutor(max_workers=a.jobs) as ex: errs = list(ex.map(_commit_task, todo))
    else:
        errs = [_commit_task(t) for t in todo]
    for e in errs:
        if e: warn('commit measures: ' + e)
    res = {}
    for k, outs in plan.items():
        rows = []
        for o in outs:
            if os.path.exists(o): rows += json.load(open(o))['rows']
        res[k] = rows
    return res


def load_systems(folders, gi=0, cm=None):
    cm = cm or {}
    raw = []
    for d in folders:
        try:
            n, t5, t7 = summ.load_pair(os.path.join(d, '05.log'), os.path.join(d, '07.log'))
        except SystemExit as e:
            warn(f'{d}: {e}'); continue
        raw.append((d, n, t5, t7))
    names = [r[1] for r in raw]
    systems = []
    for d, n, t5, t7 in raw:
        d05, (d07, meta) = summ.parse05(t5), summ.parse07(t7)
        if not d05 or not d07:
            warn(f'{d}: parsed no data in 05/07 logs'); continue
        systems.append({'name': n if names.count(n) == 1 else fname(d), 'd05': d05, 'd07': d07, 'meta': meta})
        if cm.get((gi, d)): merge_commit(systems[-1], cm[(gi, d)])
        else: warn(f'{d}: no commit measures (need hops_cif/<name>/hops_T<T>K.csv and the dumps); those columns will be empty')
    return systems


# ---------------------------------------------------------------- site deviation: ideal sites + matching
SP = ('Li', 'La', 'Zr', 'O')
SPCOL = {'Li': '#2ca02c', 'La': '#1f77b4', 'Zr': '#ff7f0e', 'O': '#d62728', 'dop': '#a020a0'}
WARN_RMS, WARN_MAX, WARN_MED, WARN_DOP = 0.35, 1.0, 1.0, 2.0     # A


def ideal_sites(cif, fold):
    """{element: (fractional sites tiled to the box (n,3), kind (n,))} for Li, La, Zr, O from the pristine CIF."""
    a, ops, atoms = ls.parse_cif(cif)
    acc = {}
    for lab, el, p in atoms:
        if el not in SP: continue
        orb = ls.orbit(ops, p)
        t = ls._tile(orb, fold)
        acc.setdefault(el, []).append((t, np.array([ls.KIND.get(len(orb), f'{len(orb)}x')] * len(t))))
    return {e: (np.vstack([x[0] for x in v]), np.concatenate([x[1] for x in v])) for e, v in acc.items()}


def pdist(fa, fs, L):
    d = fa[:, None, :] - fs[None, :, :]
    d -= np.rint(d)
    return np.linalg.norm(d * L, axis=-1)


def sd_key(p, sites, types, smooth=0.0):
    return repr((p['every'], p['tmin'], tuple(p['edges']), tuple(sorted(sites.items())), p['cif'],
                 os.path.getmtime(p['cif']), p['minframes'], tuple(sorted(types.items())), float(smooth or 0.0), 'v2'))


def md_dump(folder, T):
    """<folder>/MD-<T>/ dump: prefer a name containing '<T>K', else the only/first *.lammpstrj. None if absent."""
    hit = sorted(glob.glob(os.path.join(folder, f'MD-{T}', '*.lammpstrj')))
    pref = [h for h in hit if f'{T}K' in os.path.basename(h)]
    return (pref or hit or [None])[0]


def md_temps(folder):
    """temperatures T of every <folder>/MD-<T>/ that holds a dump, ascending."""
    out = []
    for d in glob.glob(os.path.join(folder, 'MD-*')):
        t = os.path.basename(d)[3:]
        if t.isdigit() and md_dump(folder, int(t)): out.append(int(t))
    return sorted(out)


def compute_T(folder, cfg, T, p, sites, outpath):
    """Deviation of every atom from its pure-LLZO site at temperature T -> npz."""
    traj = lio.prepare(cfg, md_dump(folder, T) or lio.dump_path(cfg, T), p['tmin'], verbose=False)
    sl = slice(None, None, max(1, p['every']))
    # dopant positions for the shells are smoothed on the FULL-resolution trajectory (--dop-smooth-ps is in ps; smoothing after
    # taking every Nth frame could shrink the window to one frame), then sampled like everything else
    dfull = lio.framework_drift(cfg, traj)
    dsm = {int(i): lio.dopant_xyz(cfg, traj, dfull, i)[sl] for t in cfg['dopant_types'] for i in np.where(traj.types == t)[0]}
    traj = dataclasses.replace(traj, steps=traj.steps[sl], time_ps=traj.time_ps[sl], L=traj.L[sl], pos=traj.pos[sl])
    F = len(traj.steps)
    L, lo = traj.L.mean(axis=0), traj.lo
    P = traj.pos - lio.framework_drift(cfg, traj)[:, None, :]
    el = np.array([cfg['types'][int(t)] for t in traj.types])
    fm = ls._w((P.mean(axis=0) - lo) / L)                          # time-mean fractional positions
    ideal = ideal_sites(p['cif'], lio.get_fold(cfg, L))
    for e in SP:
        if (el == e).any() and e not in ideal: raise ValueError(f'CIF has no {e} atoms')
    edges = np.array(p['edges'], float)
    names = lio.shell_names(edges)
    K = len(names)

    dops, cnt = [], {}
    for t in cfg['dopant_types']:
        for i in np.where(traj.types == t)[0]:
            e = cfg['types'][t]
            cnt[e] = cnt.get(e, 0) + 1
            if e not in sites: raise ValueError(f"no sublattice given for dopant {e} (use {e}=Li24d, {e}=La, ...)")
            dops.append(dict(label=f'{e}{cnt[e]}', el=e, i=int(i)))
    if not dops: raise ValueError('no dopant atoms in this folder')

    # ---- align pure LLZO on the La/Zr framework (rigid shift, same fit as 04)
    md_fw = {e: fm[el == e] for e in ('La', 'Zr') if (el == e).any()}
    t_fit, rms, mx, nfit = ls.fit_shift(md_fw, {e: ideal[e][0] for e in md_fw}, L)
    S = {e: (ls._w(ideal[e][0] + t_fit), ideal[e][1]) for e in ideal}
    warns = []
    if rms > WARN_RMS or mx > WARN_MAX:
        warns.append(f'framework fit rms {rms:.2f} / max {mx:.2f} A: alignment unreliable')

    # ---- dopants first: closest site of the stated sublattice
    taken = {e: np.zeros(len(S[e][0]), bool) for e in S}
    for d in dops:
        se, kind = sublattice(sites[d['el']])
        fs, kd = S[se]
        cand = np.where(~taken[se] & ((kd == kind) if kind else True))[0]
        if not len(cand): raise ValueError(f'no free {sites[d["el"]]} site for {d["label"]}')
        dd = pdist(fm[d['i']][None], fs[cand], L)[0]
        j = cand[np.argmin(dd)]
        taken[se][j] = True
        d.update(dev=float(dd.min()), act=fm[d['i']] * L, idl=fs[j] * L)
        if d['dev'] > WARN_DOP: warns.append(f"{d['label']} is {d['dev']:.2f} A from the nearest {sites[d['el']]} site")

    # ---- La, Zr, O: one-to-one on the free sites (time-mean positions)
    out = {}
    fwv = {}
    for e in ('La', 'Zr', 'O'):
        idx = np.where(el == e)[0]
        if not len(idx): continue
        free = np.where(~taken[e])[0]
        D = pdist(fm[idx], S[e][0][free], L)
        r, c = linear_sum_assignment(D)
        dev = np.full(len(idx), np.nan)
        dev[r] = D[r, c]
        un = np.where(np.isnan(dev))[0]
        if len(un):
            dev[un] = D[un].min(axis=1)
            warns.append(f'{e}: {len(idx)} atoms but only {len(free)} free sites ({len(un)} matched to the nearest site)')
        elif len(idx) != len(free):
            warns.append(f'{e}: {len(idx)} atoms vs {len(free)} free sites')
        if np.median(dev) > WARN_MED: warns.append(f'{e}: median deviation {np.median(dev):.2f} A (matching unreliable?)')
        fwv[e] = (idx, dev)
        out[f'act_{e}'] = fm[idx] * L
    # ---- Li: nearest free Li site, every sampled frame
    li = np.where(el == 'Li')[0]
    free = np.where(~taken['Li'])[0]
    tree = cKDTree(S['Li'][0][free] * L, boxsize=L)
    big = np.nextafter(L, 0)
    fli = ls._w((P[:, li, :] - lo) / L).reshape(-1, 3)
    dev_li = tree.query(np.minimum(fli * L, big))[0].reshape(F, len(li))
    if np.median(dev_li) > WARN_MED: warns.append(f'Li: median deviation {np.median(dev_li):.2f} A (matching unreliable?)')
    fr = np.unique(np.linspace(0, F - 1, min(F, 30)).astype(int))
    out['act_Li'] = np.minimum(ls._w((P[fr][:, li, :] - lo) / L).reshape(-1, 3) * L, big)
    for e in SP:
        if e in S: out[f'idl_{e}'] = S[e][0][np.where(~taken[e])[0]] * L

    # ---- shells (05's: distance to the dopant, far = > last edge from EVERY dopant)
    Lf = np.broadcast_to(L, (F, 3))
    # dopant position for the shells: per frame, smoothed over --dop-smooth-ps (same as 04-09); a static dopant is unchanged
    dopx = {d['label']: dsm[d['i']] for d in dops}
    dli = {d['label']: lio.dist_to(P[:, li, :], dopx[d['label']], Lf) for d in dops}
    dli['any'] = np.min([dli[d['label']] for d in dops], axis=0)
    for ref, d_this in dli.items():
        sh = lio.shell_labels(d_this, dli['any'], edges, 'all')
        v, s = [], []
        for k in range(K):
            m = sh == k
            c = m.sum(axis=0)
            ok = c >= p['minframes']
            v.append((dev_li * m).sum(axis=0)[ok] / c[ok]); s.append(np.full(ok.sum(), k))
        v.append(dev_li.mean(axis=0)); s.append(np.full(len(li), K))       # 'all'
        out[f'Li|{ref}|v'], out[f'Li|{ref}|s'] = np.concatenate(v), np.concatenate(s)
    for e, (idx, dev) in fwv.items():
        dd = {}
        for d in dops:
            x = fm[idx] - fm[d['i']]
            dd[d['label']] = np.linalg.norm((x - np.rint(x)) * L, axis=1)
        dd['any'] = np.min(list(dd.values()), axis=0)
        for ref, d_this in dd.items():
            sh = lio.shell_labels(d_this, dd['any'], edges, 'all')
            out[f'{e}|{ref}|v'] = np.concatenate([dev[sh == k] for k in range(K)] + [dev])
            out[f'{e}|{ref}|s'] = np.concatenate([np.full((sh == k).sum(), k) for k in range(K)] + [np.full(len(dev), K)])

    out.update(T=T, L=L, fit=np.array([rms, mx, nfit]), warn=np.array(warns + ['']), edges=edges,
               key=np.array(sd_key(p, sites, cfg['types'], cfg.get('dop_smooth_ps'))), dop_label=np.array([d['label'] for d in dops]),
               dop_el=np.array([d['el'] for d in dops]), dop_dev=np.array([d['dev'] for d in dops]),
               dop_act=np.array([d['act'] for d in dops]), dop_idl=np.array([d['idl'] for d in dops]))
    os.makedirs(os.path.dirname(outpath), exist_ok=True)
    np.savez_compressed(outpath, **out)


def _task(args):
    try:
        compute_T(*args)
        return None
    except (Exception, SystemExit) as e:
        return f'{args[0]} T={args[2]}: {e}'


# ---------------------------------------------------------------- site deviation: statistics
def stat(v, rng, B):
    v = np.asarray(v, float)
    n = len(v)
    r = dict(n=n, median=np.nan, lo=np.nan, hi=np.nan, mean=np.nan, std=np.nan)
    if not n: return r
    r.update(median=float(np.median(v)), mean=float(v.mean()), std=float(v.std()))
    if n >= 3:
        meds = np.median(v[rng.integers(0, n, (B, n))], axis=1)
        r['lo'], r['hi'] = (float(x) for x in np.percentile(meds, [2.5, 97.5]))
    return r


def summarize(data, names, B=1000):
    """records: median deviation (A) with 95 % bootstrap CI over atoms, per T, species, reference (any / dopant atom), shell."""
    rng = np.random.default_rng(0)
    T, K = int(data['T']), len(names)
    labels = [str(x) for x in data['dop_label']]
    recs = []
    for ref in ['any'] + labels:
        for sp in SP:
            if f'{sp}|{ref}|v' not in data: continue
            v, s = data[f'{sp}|{ref}|v'], data[f'{sp}|{ref}|s']
            for k in range(K + 1):
                recs.append(dict(T=T, species=sp, ref=ref, shell=names[k] if k < K else 'all', **stat(v[s == k], rng, B)))
        sel = data['dop_dev'] if ref == 'any' else data['dop_dev'][np.array([l == ref for l in labels])]
        recs.append(dict(T=T, species='dop', ref=ref, shell='all', **stat(sel, rng, B)))
    return recs


def dev_cell(r):
    if r is None or r['n'] == 0: return '–'
    if not np.isfinite(r['lo']): return f"{r['median']:.3f} ({r['n']})"
    return f"{r['median']:.3f} ± {(r['hi'] - r['lo']) / 2:.3f} ({r['n']})"


# ---------------------------------------------------------------- site deviation: figures
def _warn_title(fig, title, data):
    w = [str(x) for x in data['warn'] if str(x)]
    fig.suptitle(title, fontsize=8, y=0.995, va='top')
    if w: fig.text(0.5, 0.925, '[WARN] ' + '; '.join(w), ha='center', va='top', fontsize=6, color='red')


def violin_fig(data, name, recs, names, path, ylims=None):
    """ylims {panel: (lo, hi)} forces shared y limits; path=None -> no save, only returns the auto limits."""
    T = int(data['T'])
    lims = {}
    R = {(r['species'], r['ref'], r['shell']): r for r in recs}
    fig, axs = plt.subplots(1, 5, figsize=(13, 2.9), dpi=200)
    for ax, sp in zip(axs, SP + ('dop',)):
        col = SPCOL[sp]
        ax.tick_params(labelsize=6.5)
        if sp == 'dop':
            lab = [str(x) for x in data['dop_label']]
            dv = data['dop_dev']
            ax.scatter(range(len(dv)), dv, color=col, s=22, zorder=3)
            for i, v in enumerate(dv): ax.annotate(f'{v:.2f}', (i, v), xytext=(0, 4), textcoords='offset points', ha='center', fontsize=6)
            ax.set_xticks(range(len(dv))); ax.set_xticklabels(lab)
            ax.set_xlim(-0.6, len(dv) - 0.4); ax.set_title('dopant atom(s)', fontsize=8)
            ax.set_ylim(0, max(dv.max() * 1.3, 0.1))
            lims[sp] = ax.get_ylim()
            if ylims and sp in ylims: ax.set_ylim(*ylims[sp])
            continue
        ax.set_title(sp, fontsize=8)
        if f'{sp}|any|v' not in data:
            ax.text(0.5, 0.5, 'none', ha='center', transform=ax.transAxes); continue
        v, s = data[f'{sp}|any|v'], data[f'{sp}|any|s']
        ticks = []
        for k, nm in enumerate(names):
            x = v[s == k]
            n = len(x)
            ticks.append(f'{summ.shell_name(nm)}\n(n={n})')
            if n >= 5 and np.ptp(x) > 0:
                parts = ax.violinplot([x], positions=[k], widths=0.85, showextrema=False)
                for b in parts['bodies']:
                    b.set_facecolor(col); b.set_edgecolor('0.3'); b.set_alpha(0.45)
            elif n:
                ax.scatter(k + np.linspace(-0.12, 0.12, n), x, s=6, color=col, alpha=0.8)
            r = R.get((sp, 'any', nm))
            if n:
                err = None if not np.isfinite(r['lo']) else [[r['median'] - r['lo']], [r['hi'] - r['median']]]
                ax.errorbar(k, r['median'], yerr=err, fmt='o', color='k', ms=3, capsize=2, lw=1, zorder=4)
        far = R.get((sp, 'any', names[-1]))
        if far and far['n']: ax.axhline(far['median'], ls='--', color='0.5', lw=0.7, zorder=0)
        ax.set_xticks(range(len(names))); ax.set_xticklabels(ticks)
        ax.set_ylim(bottom=0)
        lims[sp] = ax.get_ylim()
        if ylims and sp in ylims: ax.set_ylim(*ylims[sp])
    axs[0].set_ylabel('deviation from pure-LLZO site (Å)', fontsize=7)
    if path is None:
        plt.close(fig)
        return lims
    _warn_title(fig, f'{name}   T = {T} K   per-site deviation by shell (distance to nearest dopant); dot = median, bar = 95% CI; dashed = >7 Å median', data)
    fig.tight_layout(rect=(0, 0, 1, 0.88))
    fig.savefig(path)
    fig.savefig(os.path.splitext(path)[0] + '.svg')
    plt.close(fig)
    return lims


PANELS = [(2, 0, 1, 'x', 'y'), (1, 0, 2, 'x', 'z'), (0, 1, 2, 'y', 'z')]       # (collapsed axis, ix, iy, xlabel, ylabel)
MS = {'Li': 14, 'La': 46, 'Zr': 34, 'O': 12}


def proj_fig(data, name, d, window, slab, path):
    T, L = int(data['T']), data['L']
    c = data['dop_act'][d]
    lab = str(data['dop_label'][d])

    def cen(x):
        y = x - c
        return y - L * np.rint(y / L)

    fig, axs = plt.subplots(1, 3, figsize=(9, 3.4), dpi=200)
    for ax, (cax, ix, iy, xl, yl) in zip(axs, PANELS):
        for sp in SP:
            if f'act_{sp}' not in data: continue
            for key, filled in ((f'idl_{sp}', False), (f'act_{sp}', True)):
                if key not in data: continue
                y = cen(data[key])
                m = (np.abs(y[:, cax]) <= slab) & (np.abs(y[:, ix]) <= window) & (np.abs(y[:, iy]) <= window)
                if filled:
                    ax.scatter(y[m, ix], y[m, iy], s=MS[sp] * (0.4 if sp == 'Li' else 1), color=SPCOL[sp], alpha=0.35 if sp == 'Li' else 0.5, lw=0, zorder=2, rasterized=True)
                else:
                    ax.scatter(y[m, ix], y[m, iy], s=MS[sp], facecolors='none', edgecolors='k', alpha=0.55, lw=0.6, zorder=3, rasterized=True)
        ax.scatter([0], [0], marker='*', s=130, color=SPCOL['dop'], alpha=0.7, lw=0, zorder=4)
        yi = cen(data['dop_idl'][d][None])[0]
        ax.scatter([yi[ix]], [yi[iy]], marker='*', s=130, facecolors='none', edgecolors='k', lw=0.8, zorder=5)
        ax.set_xlim(-window, window); ax.set_ylim(-window, window); ax.set_aspect('equal')
        ax.set_xlabel(f'{xl} (Å)', fontsize=7); ax.set_ylabel(f'{yl} (Å)', fontsize=7); ax.tick_params(labelsize=6.5)
        ax.set_title(f'±{slab:g} Å slab normal to {"xyz"[cax]}', fontsize=7)
    from matplotlib.lines import Line2D
    h = [Line2D([], [], marker='o', ls='', color=SPCOL[s], alpha=0.6, ms=5, label=f'{s} (MD)') for s in SP]
    h += [Line2D([], [], marker='*', ls='', color=SPCOL['dop'], ms=8, label=f'{lab} (MD)'),
          Line2D([], [], marker='o', ls='', mfc='none', mec='k', ms=5, label='pure-LLZO site')]
    fig.legend(handles=h, loc='lower center', ncol=7, fontsize=6, frameon=False)
    _warn_title(fig, f'{name}   T = {T} K   centred on {lab}   (Li: sampled frames; others: time-mean)', data)
    fig.tight_layout(rect=(0, 0.07, 1, 0.88))
    fig.savefig(path)
    fig.savefig(os.path.splitext(path)[0] + '.svg')
    plt.close(fig)


def _centred(data, d):
    """-> (cen function wrapping to the box centred on dopant d, L)"""
    L, c = data['L'], data['dop_act'][d]
    def cen(x):
        y = x - c
        return y - L * np.rint(y / L)
    return cen, L


def proj_full_fig(data, name, d, path):
    """Same as proj_fig but the WHOLE cell: no slab (every atom, collapsed along the axis), box centred on the dopant, periodic wrap."""
    T = int(data['T'])
    cen, L = _centred(data, d)
    lab = str(data['dop_label'][d])
    fig, axs = plt.subplots(1, 3, figsize=(11, 3.8), dpi=200)
    for ax, (cax, ix, iy, xl, yl) in zip(axs, PANELS):
        for sp in SP:
            if f'act_{sp}' not in data: continue
            for key, filled in ((f'idl_{sp}', False), (f'act_{sp}', True)):
                if key not in data: continue
                y = cen(data[key])
                if filled:
                    ax.scatter(y[:, ix], y[:, iy], s=MS[sp] * 0.5 * (0.4 if sp == 'Li' else 1), color=SPCOL[sp], alpha=0.35 if sp == 'Li' else 0.5, lw=0, zorder=2, rasterized=True)
                else:
                    ax.scatter(y[:, ix], y[:, iy], s=MS[sp] * 0.5, facecolors='none', edgecolors='k', alpha=0.55, lw=0.4, zorder=3, rasterized=True)
        ax.scatter([0], [0], marker='*', s=130, color=SPCOL['dop'], alpha=0.7, lw=0, zorder=4)
        yi = cen(data['dop_idl'][d][None])[0]
        ax.scatter([yi[ix]], [yi[iy]], marker='*', s=130, facecolors='none', edgecolors='k', lw=0.8, zorder=5)
        ax.set_xlim(-L[ix] / 2, L[ix] / 2); ax.set_ylim(-L[iy] / 2, L[iy] / 2); ax.set_aspect('equal')
        ax.set_xlabel(f'{xl} (Å)', fontsize=7); ax.set_ylabel(f'{yl} (Å)', fontsize=7); ax.tick_params(labelsize=6.5)
        ax.set_title(f'whole cell, all atoms projected along {"xyz"[cax]}', fontsize=7)
    from matplotlib.lines import Line2D
    h = [Line2D([], [], marker='o', ls='', color=SPCOL[s], alpha=0.6, ms=5, label=f'{s} (MD)') for s in SP]
    h += [Line2D([], [], marker='*', ls='', color=SPCOL['dop'], ms=8, label=f'{lab} (MD)'),
          Line2D([], [], marker='o', ls='', mfc='none', mec='k', ms=5, label='pure-LLZO site')]
    fig.legend(handles=h, loc='lower center', ncol=7, fontsize=6, frameon=False)
    _warn_title(fig, f'{name}   T = {T} K   WHOLE CELL centred on {lab}   (Li: sampled frames; others: time-mean)', data)
    fig.tight_layout(rect=(0, 0.07, 1, 0.88))
    fig.savefig(path)
    fig.savefig(os.path.splitext(path)[0] + '.svg')
    plt.close(fig)


_HTML_SIZE = {'Li': 3, 'La': 6, 'Zr': 5, 'O': 2.5}


def interactive_html(data, name, d, path):
    """Rotatable 3D whole cell (plotly), centred on dopant d. Legend click toggles a species. plotly.min.js is written once
    next to the html (offline; keep the folder together). Returns path or None if plotly is missing."""
    try:
        import plotly.graph_objects as go
    except ImportError:
        warn('plotly not installed -> no interactive html (uv pip install plotly)'); return None
    T = int(data['T'])
    cen, L = _centred(data, d)
    lab = str(data['dop_label'][d])
    fig = go.Figure()
    for sp in SP:
        for key, filled in ((f'act_{sp}', True), (f'idl_{sp}', False)):
            if key not in data: continue
            y = cen(data[key])
            fig.add_trace(go.Scatter3d(
                x=y[:, 0], y=y[:, 1], z=y[:, 2], mode='markers', name=f'{sp} {"MD" if filled else "pure-LLZO site"}',
                legendgroup=sp + ('' if filled else '_ideal'),
                marker=dict(size=_HTML_SIZE[sp] * (1 if filled else 0.8), color=SPCOL[sp], symbol='circle' if filled else 'circle-open',
                            opacity=0.55 if filled else 0.8, line=dict(width=0 if filled else 1, color='black' if not filled else SPCOL[sp])),
                customdata=np.linalg.norm(y, axis=1), hovertemplate=f'{sp}<br>%{{x:.2f}}, %{{y:.2f}}, %{{z:.2f}} Å<br>r(dopant) = %{{customdata:.2f}} Å<extra></extra>'))
    yi = cen(data['dop_idl'][d][None])[0]
    fig.add_trace(go.Scatter3d(x=[0], y=[0], z=[0], mode='markers', name=f'{lab} MD', marker=dict(size=9, color=SPCOL['dop'], symbol='diamond')))
    fig.add_trace(go.Scatter3d(x=[yi[0]], y=[yi[1]], z=[yi[2]], mode='markers', name=f'{lab} pure-LLZO site',
                               marker=dict(size=9, color='black', symbol='diamond-open')))
    rng = lambda i: [-L[i] / 2, L[i] / 2]
    fig.update_layout(title=f'{name}   T = {T} K   whole cell centred on {lab}   (Li: sampled frames; others: time-mean)',
                      scene=dict(xaxis=dict(title='x (Å)', range=rng(0)), yaxis=dict(title='y (Å)', range=rng(1)),
                                 zaxis=dict(title='z (Å)', range=rng(2)), aspectmode='data'),
                      legend=dict(itemsizing='constant'), margin=dict(l=0, r=0, t=40, b=0))
    fig.write_html(path, include_plotlyjs='directory')
    return path


def sitedev_folder(folder, a, p, sites, results):
    """npz of every T exists (computed before) -> figures, csv, records. Returns ([{T, violin, projs}], records) or None."""
    out = os.path.join(folder, 'sitedev', fname(folder))
    names = lio.shell_names(np.array(p['edges'], float))
    rows, allrecs = [], []
    loaded = []
    for T in sorted(results):
        npz = os.path.join(out, f'T{T}K.npz')
        if not os.path.exists(npz): continue
        data = dict(np.load(npz))
        loaded.append((T, data, summarize(data, names, p['boot'])))
    ylims = {}                                  # per panel (Li, La, ..., dop): one y range over all T of this folder
    for T, data, recs in loaded:
        for sp, (lo, hi) in violin_fig(data, fname(folder), recs, names, None).items():
            a0, b0 = ylims.get(sp, (lo, hi))
            ylims[sp] = (min(a0, lo), max(b0, hi))
    for T, data, recs in loaded:
        allrecs += recs
        vp = os.path.join(out, f'violin_T{T}K.png')
        violin_fig(data, fname(folder), recs, names, vp, ylims)
        pp, fp = [], []
        for d, lab in enumerate(data['dop_label']):
            pth = os.path.join(out, f'proj_T{T}K_{lab}.png')
            proj_fig(data, fname(folder), d, p['window'], p['slab'], pth)
            pp.append(pth)
            fpth = os.path.join(out, f'projfull_T{T}K_{lab}.png')
            proj_full_fig(data, fname(folder), d, fpth)
            fp.append(fpth)
            interactive_html(data, fname(folder), d, os.path.join(out, f'interactive_T{T}K_{lab}.html'))
        rows.append(dict(T=T, violin=vp, projs=pp, full=fp))
    if not rows: return None
    with open(os.path.join(out, 'site_deviation.csv'), 'w', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=list(allrecs[0].keys()))
        w.writeheader(); w.writerows(allrecs)
    return rows, allrecs


def run_sitedev(groups, a, p):
    """Compute (parallel, cached) then draw. Returns {(group index, folder): (rows, records)}."""
    jobs, todo, plan = [], [], {}
    for gi, g in enumerate(groups):
        if not g['sites']: continue
        for f in g['folders']:
            if not os.path.isdir(f): continue
            try:
                ga = argparse.Namespace(**vars(a))
                if g['elements']: ga.elements = g['elements']          # per-group element list
                cfg = lio.load_config(f, None, ga)
                temps = a.T or md_temps(f) or lio.temperatures(cfg)
            except SystemExit as e:
                warn(f'site deviation skipped for {f}: {e}'); continue
            plan[(gi, f)] = temps
            for T in temps:
                outp = os.path.join(f, 'sitedev', fname(f), f'T{T}K.npz')
                ok = False
                if os.path.exists(outp) and not a.recompute:
                    try: ok = str(np.load(outp)['key']) == sd_key(p, g['sites'], cfg['types'], cfg.get('dop_smooth_ps'))
                    except Exception: ok = False
                if not ok: todo.append((f, cfg, T, p, g['sites'], outp))
    print(f'[site deviation] {len(todo)} (folder, T) to compute, {sum(len(v) for v in plan.values()) - len(todo)} cached')
    if a.jobs > 1 and len(todo) > 1:
        from concurrent.futures import ProcessPoolExecutor
        with ProcessPoolExecutor(max_workers=a.jobs) as ex:
            errs = list(ex.map(_task, todo))
    else:
        errs = [_task(t) for t in todo]
    for e in errs:
        if e: warn('site deviation: ' + e)
    res = {}
    for (gi, f), temps in plan.items():
        r = sitedev_folder(f, a, p, groups[gi]['sites'], temps)
        if r: res[(gi, f)] = r
        else: warn(f'no site-deviation result for {f}')
    return res


# ---------------------------------------------------------------- RMSD per atom type
RMSD_EXTRA = ['#9467bd', '#8c564b', '#e377c2', '#17becf', '#bcbd22', '#7f7f7f']       # dopant elements


def rmsd_key(p, types, dump):
    return repr((p['every'], p['tmin'], tuple(sorted(types.items())), dump, os.path.getmtime(dump), 'rmsd-v2'))


def compute_rmsd(folder, cfg, T, p, outpath):
    """RMSD(frame) of every atom type vs the first frame of the window (framework drift removed) -> npz."""
    dump = md_dump(folder, T) or lio.dump_path(cfg, T)
    traj = lio.prepare(cfg, dump, p['tmin'], verbose=False)
    k = max(1, p['every'])
    traj = dataclasses.replace(traj, steps=traj.steps[::k], time_ps=traj.time_ps[::k], L=traj.L[::k], pos=traj.pos[::k])
    P = traj.pos - lio.framework_drift(cfg, traj)[:, None, :]
    L = traj.L[:, None, :]
    d = np.diff(P, axis=0)
    d -= L[1:] * np.rint(d / L[1:])                       # unwrap by minimum image (harmless if pos is already unwrapped)
    U = np.concatenate([P[:1], P[:1] + np.cumsum(d, axis=0)])
    el = np.array([cfg['types'][int(t)] for t in traj.types])
    out = dict(frame=np.arange(len(U)) * k, T=T, key=np.array(rmsd_key(p, cfg['types'], dump)))
    for e in dict.fromkeys(el.tolist()):
        m = el == e
        out['atoms_' + e] = np.sqrt(((U[:, m] - U[0, m]) ** 2).sum(axis=2))        # (frames, atoms of this type), A
    os.makedirs(os.path.dirname(outpath), exist_ok=True)
    np.savez_compressed(outpath, **out)


def _rmsd_task(args):
    try:
        compute_rmsd(*args)
        return None
    except (Exception, SystemExit) as e:
        return f'{args[0]} T={args[2]}: {e}'


def rmsd_fig(folder, temps, path):
    """Grid: row = atom type, column = T. Cell = thin line per atom + black species RMSD. x = frame (no unit), y = RMSD (A);
    y range auto-fitted to the data and the same in every panel, so atom types and temperatures compare directly."""
    out = os.path.dirname(path)
    data = [dict(np.load(os.path.join(out, f'T{T}K.npz'))) for T in temps if os.path.exists(os.path.join(out, f'T{T}K.npz'))]
    if not data: return None
    els = sorted({k[6:] for d in data for k in d if k.startswith('atoms_')})
    els = [e for e in SP if e in els] + [e for e in els if e not in SP]
    col, j = {}, 0
    for e in els:
        if e in SPCOL: col[e] = SPCOL[e]
        else: col[e] = RMSD_EXTRA[j % len(RMSD_EXTRA)]; j += 1
    nr, nc = len(els), len(data)
    fig, axs = plt.subplots(nr, nc, figsize=(max(5.0, 3.6 * nc), 1.7 * nr + 0.9), dpi=200, sharex=True, sharey='all', squeeze=False)
    for i, e in enumerate(els):
        for j, d in enumerate(data):
            ax = axs[i, j]
            ax.tick_params(labelsize=6.5)
            a = d.get('atoms_' + e)
            if a is None:
                ax.text(0.5, 0.5, 'none', ha='center', va='center', transform=ax.transAxes, fontsize=7); continue
            ax.plot(d['frame'], a, color=col[e], lw=0.4, alpha=0.45, rasterized=True)           # one thin line per atom
            ax.plot(d['frame'], np.sqrt((a ** 2).mean(axis=1)), color='k', lw=1.0)             # species RMSD
            ax.set_ylim(bottom=0)
            if i == 0: ax.set_title(f"{int(d['T'])} K", fontsize=8)
            if i == nr - 1: ax.set_xlabel('Frame', fontsize=7)
            if j == 0: ax.set_ylabel(f'{e}  RMSD (Å)', fontsize=7)
    fig.suptitle(f'{fname(folder)}   RMSD vs first frame of the window   (thin = single atom, black = species RMSD)', fontsize=8)
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    fig.savefig(path)
    fig.savefig(os.path.splitext(path)[0] + '.svg')
    plt.close(fig)
    return path


def run_rmsd(groups, a, p):
    """Compute (cached, parallel with --jobs) and draw. Returns {(group index, folder): png path}."""
    todo, plan = [], {}
    for gi, g in enumerate(groups):
        for f in g['folders']:
            if not os.path.isdir(f): continue
            try:
                ga = argparse.Namespace(**vars(a))
                if g['elements']: ga.elements = g['elements']
                cfg = lio.load_config(f, None, ga)
                temps = a.T or md_temps(f) or lio.temperatures(cfg)
            except SystemExit as e:
                warn(f'RMSD skipped for {f}: {e}'); continue
            plan[(gi, f)] = temps
            for T in temps:
                outp = os.path.join(f, 'rmsd', fname(f), f'T{T}K.npz')
                ok = False
                if os.path.exists(outp) and not a.recompute:
                    try:
                        dump = md_dump(f, T) or lio.dump_path(cfg, T)
                        ok = str(np.load(outp)['key']) == rmsd_key(p, cfg['types'], dump)
                    except Exception: ok = False
                if not ok: todo.append((f, cfg, T, p, outp))
    print(f'[rmsd] {len(todo)} (folder, T) to compute, {sum(len(v) for v in plan.values()) - len(todo)} cached')
    if a.jobs > 1 and len(todo) > 1:
        from concurrent.futures import ProcessPoolExecutor
        with ProcessPoolExecutor(max_workers=a.jobs) as ex:
            errs = list(ex.map(_rmsd_task, todo))
    else:
        errs = [_rmsd_task(t) for t in todo]
    for e in errs:
        if e: warn('rmsd: ' + e)
    res = {}
    for (gi, f), temps in plan.items():
        r = rmsd_fig(f, temps, os.path.join(f, 'rmsd', fname(f), 'rmsd.png'))
        if r: res[(gi, f)] = r
        else: warn(f'no RMSD result for {f}')
    return res


# ---------------------------------------------------------------- Li MSD / H_R by LOCAL Li density, far region only
RHO_VERSION = 'rho-v3'
RHO_LO, RHO_HI = 0.1, 1.5      # fixed relative-density range of the bins (n(<R) / cell-mean n)


def rho_key(p, cfg, dump, R, nbins, lags):
    return repr((R, nbins, tuple(lags), tuple(p['edges']), p['tmin'], cfg.get('dop_smooth_ps'), tuple(sorted(cfg['types'].items())),
                 dump, os.path.getmtime(dump), lio.DEFAULT_BLOCKS, RHO_VERSION))


def rho_rows(folder, cfg, T, p, R, nbins, lags_ps):
    """Far Li (> last shell edge from EVERY dopant at t0) binned by local Li density at t0 -> MSD and neighbourhood H_R per bin and lag."""
    traj = lio.prepare(cfg, md_dump(folder, T) or lio.dump_path(cfg, T), p['tmin'], verbose=False)
    drift = lio.framework_drift(cfg, traj)
    li = traj.pos[:, lio.li_mask(cfg, traj), :] - drift[:, None, :]
    F, N = li.shape[:2]
    dt = float(np.median(np.diff(traj.time_ps)))
    dops = [lio.dopant_xyz(cfg, traj, drift, i) for t in cfg['dopant_types'] for i in np.where(traj.types == t)[0]]
    if not dops: raise ValueError('no dopant atoms')
    dnear = np.min([lio.dist_to(li, x, traj.L) for x in dops], axis=0)
    edges = np.array(p['edges'], float)
    K, nb = len(edges) + 1, lio.DEFAULT_BLOCKS
    shl = lio.shell_labels(dnear, dnear, edges, 'all')                           # (F,N) shell of every Li at every t0
    far = shl == K - 1
    names = lio.shell_names(edges)
    lag_frames = sorted({max(1, int(round(x / dt))) for x in lags_ps})
    lag_frames = [g for g in lag_frames if F - g >= 2 * nb]
    if not lag_frames: raise ValueError('window too short for the requested lags')
    nloc = np.zeros((F, N), np.int32)
    acc = {g: [np.zeros((F - g, N), np.float32) for _ in range(3)] for g in lag_frames}     # |dr|^2, sum_S |dr_j|^2, |sum_S dr_j|^2
    for t in range(F):
        Lt = traj.L[t]
        x = np.minimum(np.mod(li[t], Lt), np.nextafter(Lt, 0))
        pr = cKDTree(x, boxsize=Lt).query_pairs(R, output_type='ndarray')
        a, b = np.concatenate([pr[:, 0], pr[:, 1]]), np.concatenate([pr[:, 1], pr[:, 0]])
        nloc[t] = np.bincount(a, minlength=N)
        for g, (D2, SN, SD) in acc.items():
            if t >= F - g: continue
            dr = li[t + g] - li[t]
            d2 = (dr * dr).sum(axis=1)
            D2[t] = d2
            SN[t] = d2 + np.bincount(a, weights=d2[b], minlength=N)
            s = dr + np.stack([np.bincount(a, weights=dr[b, c], minlength=N) for c in range(3)], axis=1)
            SD[t] = (s * s).sum(axis=1)
    nexp = (N - 1) * 4.0 / 3.0 * np.pi * R ** 3 / np.prod(traj.L, axis=1).mean()      # neighbours expected at the cell-mean Li density
    qe = np.linspace(RHO_LO, RHO_HI, nbins + 1)[1:-1] * nexp                         # FIXED edges in relative density: nbins bins over RHO_LO..RHO_HI (first/last bin open-ended)
    nbe = len(qe) + 1
    bin_all = np.searchsorted(qe, nloc, side='right')                                  # bin of every (t0, Li)
    rows = []
    for g in lag_frames:
        D2, SN, SD = acc[g]
        Fo = F - g
        for sk in range(K):
            ms = (shl[:Fo] == sk)                                                       # Li of this shell at t0
            lab = np.where(ms, bin_all[:Fo], -1)
            one = np.where(ms, 0, -1)
            cat = lambda v: np.concatenate([lio.block_counts(lab, nbe, nb, v), lio.block_counts(one, 1, nb, v)], axis=1)   # last = all of the shell
            C, S2, Sn, Sd = cat(None), cat(D2), cat(SN), cat(SD)
            ok = np.where(C.sum(axis=0)[:nbe] >= 30)[0]
            lo_, hi_ = (ok[0], ok[-1]) if len(ok) > 1 else (None, None)                 # sparsest / densest bin with enough samples
            msd, msd_se = lio.jackknife(lambda s, c: s / c, S2, C)
            rat, rat_se = lio.jackknife(lambda s, c: (s / c) / (s / c)[-1], S2, C)
            hr, hr_se = lio.jackknife(lambda n_, d_: n_ / d_, Sn, Sd)
            tb, tb_se = lio.jackknife(lambda s, c: (s[hi_] / c[hi_]) / (s[lo_] / c[lo_]), S2, C) if lo_ is not None else (np.nan, np.nan)
            for k in range(nbe + 1):
                m = (lab == k) if k < nbe else (lab >= 0)
                nl = nloc[:Fo][m]
                if not len(nl): continue
                rows.append(dict(T=int(T), tau_ps=round(float(g * dt), 6), shell=names[sk], is_far=bool(sk == K - 1),
                                 bin=('all' if k == nbe else k), n_bins=nbe, rho_r=float(R),
                                 nloc_min=int(nl.min()), nloc_max=int(nl.max()), nloc_mean=float(nl.mean()),
                                 rho_rel=float(nl.mean() / nexp), n_expected=float(nexp),
                                 n_li_per_origin=float(C[:, k].sum() / Fo), n_samples=int(C[:, k].sum()),
                                 msd_A2=float(msd[k]), msd_se=float(msd_se[k]), msd_ratio_to_all=float(rat[k]), msd_ratio_se=float(rat_se[k]),
                                 hr_local=float(hr[k]), hr_local_se=float(hr_se[k]),
                                 msd_top_over_bottom=float(tb), msd_top_over_bottom_se=float(tb_se), blocks=nb))
    return rows


def _rho_task(args):
    f, cfg, T, p, R, nbins, lags, outp, key = args
    try:
        rows = rho_rows(f, cfg, T, p, R, nbins, lags)
        os.makedirs(os.path.dirname(outp), exist_ok=True)
        with open(outp, 'w') as fh: json.dump({'key': key, 'rows': rows}, fh)
    except (Exception, SystemExit) as e:
        return f'{f} T={T}: {e}'


RHO_MIN_N = 3.0          # fewer Li of a bin per origin than this: open symbol (as 09)


def _rho_plot(lines, taus, title, path):
    """lines [(label, colour, rows of ONE line over all tau)]. Rows: MSD, H_R,loc; columns: tau. Per line: points = bins (x = mean relative
    local density), dashed = all Li of that line, open symbol = < RHO_MIN_N Li of the bin per origin."""
    fig, axs = plt.subplots(2, len(taus), figsize=(4.2 * len(taus), 6.4), dpi=150, squeeze=False)
    for ci in range(1, len(taus)): axs[1, ci].sharey(axs[1, 0])
    for ci, tau in enumerate(taus):
        tb = []
        for ri, (key, se, yl) in enumerate((('msd_A2', 'msd_se', 'MSD (Å²)'), ('hr_local', 'hr_local_se', 'H_R,loc  (1 = independent, < 1 = concerted)'))):
            ax = axs[ri, ci]
            ax.tick_params(labelsize=7)
            ax.set_xlim(RHO_LO, RHO_HI)
            for lab, col, rr in lines:
                at = [r for r in rr if abs(r['tau_ps'] - tau) < 1e-6]
                sel = sorted([r for r in at if r['bin'] != 'all'], key=lambda r: r['bin'])
                if not sel: continue
                x, y = [r['rho_rel'] for r in sel], [r[key] for r in sel]
                ax.errorbar(x, y, [r[se] if np.isfinite(r[se]) else 0 for r in sel], marker='o', ms=4, capsize=2, color=col, lw=1, label=lab)
                lo = [r for r in sel if r['n_li_per_origin'] < RHO_MIN_N]
                if lo: ax.plot([r['rho_rel'] for r in lo], [r[key] for r in lo], ls='', marker='o', ms=5, mfc='white', mec=col, zorder=5)
                al = [r for r in at if r['bin'] == 'all']
                if al: ax.axhline(al[0][key], color=col, ls='--', lw=0.6, alpha=0.7)
                if ri == 0 and np.isfinite(sel[0]['msd_top_over_bottom']):
                    tb.append(f"{lab}: top/bottom {sel[0]['msd_top_over_bottom']:.2f} ± {sel[0]['msd_top_over_bottom_se']:.2f}")
            if ri == 1: ax.axhline(1, color='0.5', lw=0.8); ax.set_xlabel('local Li density at t0, n(<R) / n(cell mean)', fontsize=8)
            if ci == 0: ax.set_ylabel(yl, fontsize=8)
            if ri == 0: ax.set_title(f'tau = {tau:g} ps\n' + '\n'.join(tb), fontsize=8)
    h, l = axs[0, 0].get_legend_handles_labels()
    for ax in axs.ravel():
        if not h: h, l = ax.get_legend_handles_labels()
    fig.legend(h, l, loc='lower center', ncol=max(len(l), 1), fontsize=7, frameon=False)
    fig.suptitle(title, fontsize=7)
    fig.tight_layout(rect=(0, 0.04, 1, 0.95))
    fig.savefig(path)
    fig.savefig(os.path.splitext(path)[0] + '.svg')
    plt.close(fig)


def rho_fig(folder, rows, path):
    """Far Li (> last shell edge from every dopant): one line per T."""
    far = [r for r in rows if r['is_far']]
    Ts = sorted({r['T'] for r in far})
    cols = [plt.cm.viridis(i / max(len(Ts) - 1, 1) * 0.85) for i in range(len(Ts))]
    R = rows[0]['rho_r']
    _rho_plot([(f'{T} K', c, [r for r in far if r['T'] == T]) for T, c in zip(Ts, cols)], sorted({r['tau_ps'] for r in far}),
              f'{fname(folder)}   Li far from every dopant (> last shell edge) split by local Li density (neighbours within {R:g} Å)\n'
              'error bars = block jackknife; dashed = all far Li; open = < 3 Li per origin; top/bottom = MSD(densest bin) / MSD(sparsest bin)', path)


def rho_shell_fig(folder, rows, T, path):
    """One T: one line per distance shell, the SAME density bins (edges from the far Li) for every shell."""
    rt = [r for r in rows if r['T'] == T]
    names = list(dict.fromkeys(r['shell'] for r in rt))
    cols = [plt.cm.viridis(k / max(len(names) - 2, 1) * 0.85) for k in range(len(names) - 1)] + ['k']
    R = rows[0]['rho_r']
    _rho_plot([(nm, c, [r for r in rt if r['shell'] == nm]) for nm, c in zip(names, cols)], sorted({r['tau_ps'] for r in rt}),
              f'{fname(folder)}   T = {T} K   Li of every distance shell (to the nearest dopant at t0) split by local Li density (neighbours within {R:g} Å), '
              'bins fixed by the far Li\nerror bars = block jackknife; dashed = all Li of the shell; open = < 3 Li of that bin per origin (noisy)\n'
              'top/bottom = MSD(densest) / MSD(sparsest) bin with >= 30 samples', path)


def run_rho(groups, a, p):
    """Compute (cached, parallel with --jobs), write csv + figures. Returns {(group index, folder): [png paths]}."""
    todo, plan = [], {}
    for gi, g in enumerate(groups):
        for f in g['folders']:
            if not os.path.isdir(f): continue
            try:
                ga = argparse.Namespace(**vars(a))
                if g['elements']: ga.elements = g['elements']
                cfg = lio.load_config(f, None, ga)
                temps = a.T or md_temps(f) or lio.temperatures(cfg)
                if not cfg['dopant_types']: warn(f'local density bins skipped for {f}: no dopant atoms'); continue
            except SystemExit as e:
                warn(f'local density bins skipped for {f}: {e}'); continue
            for T in temps:
                dump = md_dump(f, T) or lio.dump_path(cfg, T)
                key = rho_key(p, cfg, dump, a.rho_r, a.rho_bins, a.rho_lags)
                outp = os.path.join(f, 'rhobin', fname(f), f'rho_T{T}K.json')
                plan.setdefault((gi, f), []).append(outp)
                ok = False
                if os.path.exists(outp) and not a.recompute:
                    try: ok = json.load(open(outp))['key'] == key
                    except Exception: ok = False
                if not ok: todo.append((f, cfg, T, p, a.rho_r, a.rho_bins, a.rho_lags, outp, key))
    print(f'[local density bins] {len(todo)} (folder, T) to compute, {sum(len(v) for v in plan.values()) - len(todo)} cached')
    if a.jobs > 1 and len(todo) > 1:
        from concurrent.futures import ProcessPoolExecutor
        with ProcessPoolExecutor(max_workers=a.jobs) as ex: errs = list(ex.map(_rho_task, todo))
    else:
        errs = [_rho_task(t) for t in todo]
    for e in errs:
        if e: warn('local density bins: ' + e)
    res = {}
    for (gi, f), outs in plan.items():
        rows = []
        for o in outs:
            if os.path.exists(o): rows += json.load(open(o))['rows']
        if not rows: warn(f'no local-density result for {f}'); continue
        d = os.path.join(f, 'rhobin', fname(f))
        with open(os.path.join(d, 'rho_bins.csv'), 'w', newline='') as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
            w.writeheader(); w.writerows(rows)
        rho_fig(f, rows, os.path.join(d, 'rho_bins.png'))
        res[(gi, f)] = [os.path.join(d, 'rho_bins.png')]
        for T in sorted({r['T'] for r in rows}):
            rho_shell_fig(f, rows, T, os.path.join(d, f'rho_shells_T{T}K.png'))
            res[(gi, f)].append(os.path.join(d, f'rho_shells_T{T}K.png'))
    return res


# ---------------------------------------------------------------- shell Li density vs MSD / H_R, pooled over folders, + non-dopant control
SP_VERSION = 'shellpool-v1'
SP_MARK = 'osv^D<>pPhX*'


def sp_key(p, cfg, dump, lags, nctrl, seed, ref):
    return repr((tuple(lags), tuple(p['edges']), p['tmin'], cfg.get('dop_smooth_ps'), tuple(sorted(cfg['types'].items())), dump,
                 os.path.getmtime(dump), nctrl, seed, ref, lio.DEFAULT_BLOCKS, SP_VERSION))


def _mi_dist(a, b, L):
    d = np.asarray(a, float) - np.asarray(b, float)
    return float(np.linalg.norm(d - L * np.rint(d / L)))


def _ctrl_vol(c, dopm, L, edges, n=48):
    """Fraction of the box volume in each shell of the control point c (grid, minimum image); grid points within the last edge
    of any real dopant are excluded (the Li there are left out of the control)."""
    g = [(np.arange(n) + 0.5) / n * L[i] for i in range(3)]
    P = np.stack(np.meshgrid(*g, indexing='ij'), axis=-1).reshape(-1, 3)

    def mi(x):
        d = P - np.asarray(x, float)
        d -= L * np.rint(d / L)
        return np.linalg.norm(d, axis=1)

    sh = np.searchsorted(edges, mi(c))
    dr = np.min([mi(x) for x in dopm], axis=0)
    sh = np.where(dr <= edges[-1], -1, sh)
    return np.array([(sh == k).mean() for k in range(len(edges) + 1)])


def shellpool_rows(folder, cfg, T, p, lags_ps, nctrl, seed, ref):
    """Per shell and lag: Li density ratio, MSD, MSD/far (05) and H_R (09), for the real dopants ('dopant') and for non-dopant
    reference points ('control'). Uses 05's shell_msd and 09's shell_sums, so the numbers are theirs."""
    import importlib
    m05, h09 = importlib.import_module('05_msd_by_shell'), importlib.import_module('09_shell_haven')
    traj = lio.prepare(cfg, md_dump(folder, T) or lio.dump_path(cfg, T), p['tmin'], verbose=False)
    drift = lio.framework_drift(cfg, traj)
    li = traj.pos[:, lio.li_mask(cfg, traj), :] - drift[:, None, :]
    F, N = li.shape[:2]
    L = traj.L.mean(axis=0)
    Lf = np.broadcast_to(L, (F, 3))
    dt = float(np.median(np.diff(traj.time_ps)))
    dops = [lio.dopant_xyz(cfg, traj, drift, i) for t in cfg['dopant_types'] for i in np.where(traj.types == t)[0]]
    if not dops: raise ValueError('no dopant atoms')
    edges = np.array(p['edges'], float)
    names = lio.shell_names(edges)
    K, nb = len(names), lio.DEFAULT_BLOCKS
    pairs = {}
    for x in lags_ps: pairs.setdefault(max(1, int(round(x / dt))), float(x))
    pairs = {g: x for g, x in pairs.items() if F - g >= 2 * nb}
    if not pairs: raise ValueError('window too short for the requested lags')
    lag_frames = sorted(pairs)
    D = {g: li[g:] - li[:-g] for g in lag_frames}
    frb = np.bincount(lio.block_id(np.arange(F), F, nb), minlength=nb)[:, None].astype(float)
    dnear = np.min([lio.dist_to(li, x, Lf) for x in dops], axis=0)

    def block(kind, refname, sh, vf):
        expect = N * vf
        C = lio.block_counts(sh, K, nb)
        dens, dens_se = lio.jackknife(lambda c_, f_: c_ / f_ / expect, C, frb)
        res = m05.shell_msd(li, sh, K, lag_frames, nb, 'start')
        out = []
        for g in lag_frames:
            msd, se, rat, rse, _ = res[g]
            NUM, DEN, CNT, ORG = h09.shell_sums(D[g], sh[:-g], K, nb)
            H, Hse = lio.jackknife(lambda n_, d_: n_ / d_, NUM, DEN)
            nm = CNT.sum(0) / ORG.sum()
            for k in range(K):
                out.append(dict(folder=fname(folder), T=int(T), kind=kind, ref=refname, tau_req=pairs[g], tau_ps=float(g * dt),
                                shell=names[k], is_far=bool(k == K - 1), n_li=float(nm[k]), density_ratio=float(dens[k]),
                                density_se=float(dens_se[k]), msd_A2=float(msd[k]), msd_se=float(se[k]),
                                ratio_to_far=float(rat[k]), ratio_se=float(rse[k]), H_R=float(H[k]), H_R_se=float(Hse[k])))
        return out

    rows = block('dopant', 'any', lio.shell_labels(dnear, dnear, edges, 'all'),
                 lio.shell_volume_fractions({f'd{i}': x.mean(axis=0) for i, x in enumerate(dops)}, L, edges, 'all')['any'])
    dopm = [x.mean(axis=0) for x in dops]
    rng = np.random.default_rng(seed)
    cand = li[0][rng.permutation(N)] if ref == 'li' else rng.random((4000, 3)) * L
    pts = []
    for c in cand:
        if all(_mi_dist(c, d_, L) >= edges[-1] for d_ in dopm): pts.append(np.asarray(c, float))
        if len(pts) >= nctrl: break
    if not pts: warn(f'{folder} T={T}: no control point >= {edges[-1]:g} A from every dopant (cell too small?): control skipped')
    for i, c in enumerate(pts):
        dc = lio.dist_to(li, np.broadcast_to(c, (F, 3)), Lf)
        sh = np.where(dnear <= edges[-1], -1, lio.shell_labels(dc, dc, edges, 'self'))        # no real dopant inside a control shell
        rows += block('control', f'c{i}', sh, _ctrl_vol(c, dopm, L, edges))
    return rows


def _sp_task(args):
    f, cfg, T, p, lags, nctrl, seed, ref, outp, key = args
    try:
        rows = shellpool_rows(f, cfg, T, p, lags, nctrl, seed, ref)
        os.makedirs(os.path.dirname(outp), exist_ok=True)
        with open(outp, 'w') as fh: json.dump({'key': key, 'rows': rows}, fh)
    except (Exception, SystemExit) as e:
        return f'{f} T={T}: {e}'


def run_shellpool(groups, a, p):
    """Compute (cached in <folder>/shellpool/, parallel with --jobs). Returns {(group index, folder): [rows]}."""
    todo, plan = [], {}
    for gi, g in enumerate(groups):
        for f in g['folders']:
            if not os.path.isdir(f): continue
            try:
                ga = argparse.Namespace(**vars(a))
                if g['elements']: ga.elements = g['elements']
                cfg = lio.load_config(f, None, ga)
                temps = a.T or md_temps(f) or lio.temperatures(cfg)
                if not cfg['dopant_types']: warn(f'shell density pool skipped for {f}: no dopant atoms'); continue
            except SystemExit as e:
                warn(f'shell density pool skipped for {f}: {e}'); continue
            for T in temps:
                dump = md_dump(f, T) or lio.dump_path(cfg, T)
                key = sp_key(p, cfg, dump, a.pool_lags, a.ctrl_n, a.ctrl_seed, a.ctrl_ref)
                outp = os.path.join(f, 'shellpool', fname(f), f'shellpool_T{T}K.json')
                plan.setdefault((gi, f), []).append(outp)
                ok = False
                if os.path.exists(outp) and not a.recompute:
                    try: ok = json.load(open(outp))['key'] == key
                    except Exception: ok = False
                if not ok: todo.append((f, cfg, T, p, a.pool_lags, a.ctrl_n, a.ctrl_seed, a.ctrl_ref, outp, key))
    print(f'[shell density pool] {len(todo)} (folder, T) to compute, {sum(len(v) for v in plan.values()) - len(todo)} cached')
    if a.jobs > 1 and len(todo) > 1:
        from concurrent.futures import ProcessPoolExecutor
        with ProcessPoolExecutor(max_workers=a.jobs) as ex: errs = list(ex.map(_sp_task, todo))
    else:
        errs = [_sp_task(t) for t in todo]
    for e in errs:
        if e: warn('shell density pool: ' + e)
    res = {}
    for k, outs in plan.items():
        rows = []
        for o in outs:
            if os.path.exists(o): rows += json.load(open(o))['rows']
        if rows: res[k] = rows
        else: warn(f'no shell density result for {k[1]}')
    return res


SP_METRICS = [('msd_A2', 'msd_se', 'MSD (Å²)'), ('ratio_to_far', 'ratio_se', 'MSD(shell) / MSD(far)'),
              ('H_R', 'H_R_se', 'H_R  (1 = independent, < 1 = concerted)')]
SP_TEXT = ('Each dot = one folder in one distance shell (colour = shell, marker = folder). x = Li density in the shell relative to a uniform '
           'density (1 = cell mean); y = MSD, MSD / MSD(far) or intra-shell H_R (09). Pale dot = fewer Li in the shell than the minimum '
           'set by --pool-min-n (H_R tends to 1 by construction). The number in each panel is the Spearman correlation of x and y '
           'across folders within one shell (n = folders, needs 4): if it is clearly non-zero, density predicts MSD / H_R at fixed distance '
           'to the dopant; if only the trend between shells exists, density is confounded with distance. Diamonds = the same shells '
           'around Li-site reference points that are not dopants (mean over points): if they show the same pattern, it is '
           'lattice geometry, not the dopant.')


def make_pool(sp, outdir, min_n):
    """Pool the rows of every folder, write csv + one figure per T. Returns {'text', 'figs'} or None."""
    from scipy.stats import pearsonr, spearmanr
    from matplotlib.lines import Line2D
    rows, seen = [], set()
    for (gi, f), rr in sp.items():
        if os.path.abspath(f) in seen: continue
        seen.add(os.path.abspath(f)); rows += rr
    if not rows: return None
    os.makedirs(outdir, exist_ok=True)
    with open(os.path.join(outdir, 'shell_pool.csv'), 'w', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)

    def corr(xy):
        xy = [(x, y) for x, y in xy if np.isfinite(x) and np.isfinite(y)]
        if len(xy) < 4: return len(xy), np.nan, np.nan, np.nan
        x, y = np.array(xy).T
        if np.ptp(x) == 0 or np.ptp(y) == 0: return len(xy), np.nan, np.nan, np.nan
        r, pv = spearmanr(x, y)
        return len(xy), float(r), float(pv), float(pearsonr(x, y)[0])

    cr, figs = [], []
    for T in sorted({r['T'] for r in rows}):
        rt = [r for r in rows if r['T'] == T]
        real, ctrl = [r for r in rt if r['kind'] == 'dopant'], [r for r in rt if r['kind'] == 'control']
        taus = sorted({r['tau_req'] for r in rt})
        names = list(dict.fromkeys(r['shell'] for r in rt))
        K = len(names)
        cols = [plt.cm.viridis(k / max(K - 2, 1) * 0.85) for k in range(K - 1)] + ['k']
        folders = sorted({r['folder'] for r in real}, key=nat)
        mk = {fo: SP_MARK[i % len(SP_MARK)] for i, fo in enumerate(folders)}
        fig, axs = plt.subplots(3, len(taus), figsize=(4.4 * len(taus), 9.6), dpi=150, squeeze=False)
        for ci, tau in enumerate(taus):
            for ri, (key, se, yl) in enumerate(SP_METRICS):
                ax = axs[ri, ci]
                ax.tick_params(labelsize=7)
                lines, allxy = [], []
                for k, nm in enumerate(names):
                    if key == 'ratio_to_far' and k == K - 1: continue                 # far / far = 1
                    sel = [r for r in real if r['tau_req'] == tau and r['shell'] == nm]
                    for r in sel:
                        ax.plot(r['density_ratio'], r[key], ls='', marker=mk[r['folder']], ms=4, color=cols[k],
                                alpha=1.0 if r['n_li'] >= min_n else 0.3)
                    use = [(r['density_ratio'], r[key]) for r in sel if r['n_li'] >= min_n]
                    allxy += use
                    n, rho, pv, pr = corr(use)
                    cr.append(dict(T=T, tau_ps=tau, shell=nm, metric=key, n=n, spearman=rho, p_spearman=pv, pearson=pr))
                    lines.append(f'{nm}: rho = ' + (f'{rho:+.2f}' if np.isfinite(rho) else 'n/a') + f' (n={n})')
                    cs = [r for r in ctrl if r['tau_req'] == tau and r['shell'] == nm and np.isfinite(r[key]) and np.isfinite(r['density_ratio'])]
                    if cs:
                        cx, cy = np.array([r['density_ratio'] for r in cs]), np.array([r[key] for r in cs])
                        ax.plot(cx.mean(), cy.mean(), ls='', marker='D', ms=6, mfc='white', mec=cols[k], mew=1.3, zorder=5)
                n, rho, pv, pr = corr(allxy)
                cr.append(dict(T=T, tau_ps=tau, shell='all shells', metric=key, n=n, spearman=rho, p_spearman=pv, pearson=pr))
                lines.append('all shells: rho = ' + (f'{rho:+.2f}' if np.isfinite(rho) else 'n/a') + f' (n={n})')
                ax.text(0.02, 0.98, '\n'.join(lines), transform=ax.transAxes, va='top', fontsize=5.5, family='monospace')
                if ri == 2: ax.axhline(1, color='0.5', lw=0.8)
                if ri == 2: ax.set_xlabel('Li density in shell / uniform', fontsize=8)
                if ci == 0: ax.set_ylabel(yl, fontsize=8)
                if ri == 0: ax.set_title(f'tau = {tau:g} ps', fontsize=9)
        h = [Line2D([], [], marker='o', ls='', color=cols[k], label=f'shell {names[k]}') for k in range(K)]
        h += [Line2D([], [], marker=mk[fo], ls='', color='0.4', label=fo) for fo in folders]
        h += [Line2D([], [], marker='D', ls='', mfc='white', mec='0.3', label='control: non-dopant Li-site points (mean)'),
              Line2D([], [], marker='o', ls='', color='0.4', alpha=0.3, label=f'pale: < {min_n:g} Li in the shell')]
        ncol = min(5, len(h))
        nrow = -(-len(h) // ncol)
        fig.legend(handles=h, loc='lower center', ncol=ncol, fontsize=6, frameon=False)
        fig.suptitle(f'T = {T} K   Li density in a distance shell vs MSD and H_R   (one dot per folder and shell, no error bars; rho = Spearman '
                     f'across folders, within a shell)', fontsize=8)
        fig.tight_layout(rect=(0, 0.025 * nrow + 0.01, 1, 0.97))
        pth = os.path.join(outdir, f'pool_T{T}K.png')
        fig.savefig(pth); fig.savefig(os.path.splitext(pth)[0] + '.svg')
        plt.close(fig)
        figs.append(pth)
    with open(os.path.join(outdir, 'shell_corr.csv'), 'w', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=list(cr[0].keys())); w.writeheader(); w.writerows(cr)
    for T in sorted({c['T'] for c in cr}):
        for key in ('ratio_to_far', 'H_R'):
            for tau in sorted({c['tau_ps'] for c in cr}):
                xs = [c for c in cr if c['T'] == T and c['metric'] == key and c['tau_ps'] == tau]
                inn = [c['spearman'] for c in xs if c['shell'] != 'all shells' and np.isfinite(c['spearman'])]
                al = [c['spearman'] for c in xs if c['shell'] == 'all shells']
                print(f'[shell pool] T={T} {key} tau={tau:g}: all-shell rho = ' + (f'{al[0]:+.2f}' if al and np.isfinite(al[0]) else 'n/a')
                      + '; within-shell rho = ' + (', '.join(f'{v:+.2f}' for v in inn) if inn else 'n/a (need >= 4 folders per shell)'))
    print(f'[shell pool] {outdir}/shell_pool.csv shell_corr.csv ' + ' '.join(os.path.basename(f) for f in figs))
    return dict(text=SP_TEXT, figs=figs)


# ---------------------------------------------------------------- radius of influence of a dopant: continuous MSD(r) profile, no shells
RAD_VERSION = 'radius-v1'
RAD_MODELS = ('exp', 'gauss', 'sigmoid')
RAD_COL = {'exp': '#1f77b4', 'gauss': '#2ca02c', 'sigmoid': '#d62728'}
RAD_SUM_MODELS = ('sigmoid',)                   # models drawn in the summary figure (exp / gauss fail on real data: B hits its bound)
RAD_A_MIN = 0.1                                 # radius-vs-concentration plot: r95 not drawn below this depth A
RAD_DELTAS = (0.05, 0.10)                       # r95 / r90: radius where the fitted profile is within 5 / 10 % (of the plateau) of the plateau
RAD_COLS = ['level', 'name', 'T', 'tau_ps', 'model', 'best', 'n_folders', 'n_dopants', 'A', 'A_lo', 'A_hi', 'lam', 'r0', 'w', 'B', 'B_lo', 'B_hi',
            'r50', 'r50_lo', 'r50_hi', 'r90', 'r90_lo', 'r90_hi', 'r95', 'r95_lo', 'r95_hi', 'rmse', 'aic',
            'r_noise', 'r_noise_lo', 'r_noise_hi', 'r_eff_null', 'ell_A', 'ell_A_lo', 'ell_A_hi', 'V_A3', 'V_A3_lo', 'V_A3_hi', 'r_eff', 'r_eff_lo', 'r_eff_hi',
            'r_max_used', 'dopant_dopant_min_A', 'half_box_A', 'r_p99_A']
_trapz = getattr(np, 'trapezoid', None) or np.trapz


def _rad_dist(li, x, L):
    """Minimum-image distance of every Li (F, N, 3) to a point that moves with the frame (F, 3); L = (3,) cell. -> (F, N)"""
    d = li - x[:, None, :]
    d -= L * np.rint(d / L)
    return np.sqrt((d * d).sum(axis=2))


def _rad_ksum(r, m, rg, bw):
    """Gaussian-kernel sums at the grid rg (width bw): row 0 = sum w, row 1 = sum w*m, row 2 = sum w^2 (kernel n_eff = row0^2 / row2).
    Mean MSD at r = row1 / row0. Kernel regression, so no bin edges."""
    out = np.zeros((3, len(rg)))
    ok = (r > rg[0] - 4 * bw) & (r < rg[-1] + 4 * bw)
    r, m = r[ok], m[ok]
    for i in range(0, len(r), 50000):
        w = np.exp(-0.5 * ((r[i:i + 50000, None] - rg[None, :]) / bw) ** 2)
        out[0] += w.sum(0); out[1] += m[i:i + 50000] @ w; out[2] += (w * w).sum(0)
    return out


def radius_core(li, dops, L, dt, lags_ps, rg, bw, rfar, nctrl, seed, ref, bid, nb):
    """li (F, N, 3) Li (unwrapped, framework drift removed); dops = list of (F, 3) dopant positions; L (3,); bid = time-block id of every frame.
    Per lag: kernel sums of the Li MSD vs distance to the NEAREST dopant at t0 ('start' rule, as 05), per time block; the far-Li MSD sums
    (Li farther than rfar from every dopant at t0, the normaliser); the same kernel sums around nctrl non-dopant control points (only Li that
    are far from every real dopant, so the control profile is 1 on average by construction)."""
    F, N = li.shape[:2]
    dnear = np.min([_rad_dist(li, x, L) for x in dops], axis=0)                      # (F, N)
    pairs = {}
    for x in lags_ps: pairs.setdefault(max(1, int(round(x / dt))), float(x))
    gone = [x for g, x in pairs.items() if F - g < 2 * nb]
    if gone: warn(f'radius: window too short for tau = {gone} ps, skipped')
    pairs = {g: x for g, x in pairs.items() if F - g >= 2 * nb}
    if not pairs: raise ValueError('window too short for the requested lags')
    dopm = [x.mean(axis=0) for x in dops]
    rng = np.random.default_rng(seed)
    cand = li[0][rng.permutation(N)] if ref == 'li' else rng.random((4000, 3)) * L
    pts = []
    for c in cand:
        if all(_mi_dist(c, d_, L) >= rfar for d_ in dopm): pts.append(np.asarray(c, float))
        if len(pts) >= nctrl: break
    if not pts: warn(f'radius: no control point >= {rfar:g} A from every dopant (cell too small?): control skipped')
    MM = {g: ((li[g:] - li[:-g]) ** 2).sum(axis=2) for g in pairs}                   # |dr|^2 per origin frame and Li, (F-g, N)
    blk = {g: [np.where(bid[:F - g] == b)[0] for b in range(nb)] for g in pairs}
    G = len(rg)
    res = {g: dict(tau_req=pairs[g], tau_ps=float(g * dt), S=np.zeros((nb, 3, G)), far=np.zeros((nb, 2)), ctrl=np.zeros((len(pts), nb, 3, G)))
           for g in pairs}
    for g in pairs:
        for b, fr in enumerate(blk[g]):
            r, m = dnear[fr], MM[g][fr]
            res[g]['S'][b] = _rad_ksum(r.ravel(), m.ravel(), rg, bw)
            far = r > rfar
            res[g]['far'][b] = (far.sum(), m[far].sum())
    for i, c in enumerate(pts):
        dc = _rad_dist(li, np.broadcast_to(c, (F, 3)), L)
        for g in pairs:
            for b, fr in enumerate(blk[g]):
                el = dnear[fr] > rfar                                                    # no real dopant within rfar of a control Li
                res[g]['ctrl'][i, b] = _rad_ksum(dc[fr][el], MM[g][fr][el], rg, bw)
    dd = [_mi_dist(a_, b_, L) for i, a_ in enumerate(dopm) for b_ in dopm[i + 1:]]
    meta = dict(n_dop=len(dops), dd_min=float(min(dd)) if dd else float('nan'), half_box=float(np.min(L) / 2),
                r_p99=float(np.percentile(dnear, 99)), n_ctrl=len(pts))
    return meta, [res[g] for g in sorted(res)]


def radius_rows(folder, cfg, T, p, lags_ps, rg, bw, rfar, nctrl, seed, ref):
    traj = lio.prepare(cfg, md_dump(folder, T) or lio.dump_path(cfg, T), p['tmin'], verbose=False)
    drift = lio.framework_drift(cfg, traj)
    li = traj.pos[:, lio.li_mask(cfg, traj), :] - drift[:, None, :]
    F = li.shape[0]
    L = np.asarray(traj.L.mean(axis=0), float)
    dt = float(np.median(np.diff(traj.time_ps)))
    dops = [lio.dopant_xyz(cfg, traj, drift, i) for t in cfg['dopant_types'] for i in np.where(traj.types == t)[0]]
    if not dops: raise ValueError('no dopant atoms')
    nb = lio.DEFAULT_BLOCKS
    bid = np.asarray(lio.block_id(np.arange(F), F, nb))
    meta, lags = radius_core(li, dops, L, dt, lags_ps, rg, bw, rfar, nctrl, seed, ref, bid, nb)
    label = '+'.join(sorted({str(cfg['types'].get(int(t), t)) for t in cfg['dopant_types']}))
    return dict(folder=fname(folder), label=label, T=int(T), rg=rg.tolist(), bw=bw, rfar=rfar, meta=meta,
                lags=[{k: (v.tolist() if isinstance(v, np.ndarray) else v) for k, v in lg.items()} for lg in lags])


def rad_key(p, cfg, dump, lags, rg, bw, rfar, nctrl, seed, ref):
    return repr((tuple(lags), tuple(np.round(rg, 6)), bw, rfar, p['tmin'], cfg.get('dop_smooth_ps'), tuple(sorted(cfg['types'].items())), dump,
                 os.path.getmtime(dump), nctrl, seed, ref, lio.DEFAULT_BLOCKS, RAD_VERSION))


def _rad_task(args):
    f, cfg, T, p, lags, rg, bw, rfar, nctrl, seed, ref, outp, key = args
    try:
        out = radius_rows(f, cfg, T, p, lags, rg, bw, rfar, nctrl, seed, ref)
        out['key'] = key
        os.makedirs(os.path.dirname(outp), exist_ok=True)
        with open(outp, 'w') as fh: json.dump(out, fh)
    except (Exception, SystemExit) as e:
        return f'{f} T={T}: {e}'


def run_radius(groups, a, p):
    """Compute the per-(folder, T) kernel sums (cached in <folder>/radius/, parallel with --jobs). Returns {(group index, folder): [results per T]}."""
    rg = np.arange(a.rad_step, a.rad_rmax + 1e-9, a.rad_step)
    rfar = float(a.rad_far) if a.rad_far else float(p['edges'][-1])
    todo, plan = [], {}
    for gi, g in enumerate(groups):
        for f in g['folders']:
            if not os.path.isdir(f): continue
            try:
                ga = argparse.Namespace(**vars(a))
                if g['elements']: ga.elements = g['elements']
                cfg = lio.load_config(f, None, ga)
                temps = a.T or md_temps(f) or lio.temperatures(cfg)
                if not cfg['dopant_types']: warn(f'radius skipped for {f}: no dopant atoms'); continue
            except SystemExit as e:
                warn(f'radius skipped for {f}: {e}'); continue
            for T in temps:
                dump = md_dump(f, T) or lio.dump_path(cfg, T)
                key = rad_key(p, cfg, dump, a.rad_lags, rg, a.rad_bw, rfar, a.ctrl_n, a.ctrl_seed, a.ctrl_ref)
                outp = os.path.join(f, 'radius', fname(f), f'radius_T{T}K.json')
                plan.setdefault((gi, f), []).append(outp)
                ok = False
                if os.path.exists(outp) and not a.recompute:
                    try: ok = json.load(open(outp))['key'] == key
                    except Exception: ok = False
                if not ok: todo.append((f, cfg, T, p, a.rad_lags, rg, a.rad_bw, rfar, a.ctrl_n, a.ctrl_seed, a.ctrl_ref, outp, key))
    print(f'[radius] {len(todo)} (folder, T) to compute, {sum(len(v) for v in plan.values()) - len(todo)} cached')
    if a.jobs > 1 and len(todo) > 1:
        from concurrent.futures import ProcessPoolExecutor
        with ProcessPoolExecutor(max_workers=a.jobs) as ex: errs = list(ex.map(_rad_task, todo))
    else:
        errs = [_rad_task(t) for t in todo]
    for e in errs:
        if e: warn('radius: ' + e)
    res = {}
    for k, outs in plan.items():
        rr = [json.load(open(o)) for o in outs if os.path.exists(o)]
        if rr: res[k] = rr
        else: warn(f'no radius result for {k[1]}')
    return res


# ---- fits. y = MSD(r) / MSD(far) with a free plateau B (B > 1: still rising beyond the far edge, i.e. the influence reaches past it)
def _rf_exp(r, A, lam, B): return B * (1 - A * np.exp(-r / lam))
def _rf_gauss(r, A, lam, B): return B * (1 - A * np.exp(-(r / lam) ** 2))
def _rf_sig(r, A, r0, w, B): return B * (1 - A / (1 + np.exp((r - r0) / w)))


RAD_FIT = {'exp': dict(f=_rf_exp, names=('A', 'lam', 'B'), lo=(0, 0.1, 0.5), hi=(1.5, 40, 1.5), p0=lambda A: (A, 2.0, 1.0)),
           'gauss': dict(f=_rf_gauss, names=('A', 'lam', 'B'), lo=(0, 0.1, 0.5), hi=(1.5, 40, 1.5), p0=lambda A: (A, 3.0, 1.0)),
           'sigmoid': dict(f=_rf_sig, names=('A', 'r0', 'w', 'B'), lo=(0, 0, 0.05, 0.5), hi=(1.5, 40, 10, 1.5), p0=lambda A: (A, 3.0, 0.7, 1.0))}


def _rad_radii(model, par):
    """r90 / r95 = radius where the fitted deficit has fallen to 10 / 5 % of the plateau; r50 = half recovery (the shape-free scale).
    exp: deficit A exp(-r/lam); gauss: A exp(-(r/lam)^2); sigmoid: A / (1 + exp((r - r0)/w))."""
    A, out = par['A'], {}
    for d in RAD_DELTAS:
        k = 'r%d' % round(100 * (1 - d))
        if A <= d: out[k] = 0.0
        elif model == 'exp': out[k] = par['lam'] * np.log(A / d)
        elif model == 'gauss': out[k] = par['lam'] * np.sqrt(np.log(A / d))
        else: out[k] = max(0.0, par['r0'] + par['w'] * np.log(A / d - 1))
    out['r50'] = 0.0 if A <= 0 else {'exp': lambda: par['lam'] * np.log(2), 'gauss': lambda: par['lam'] * np.sqrt(np.log(2)),
                                     'sigmoid': lambda: par['r0']}[model]()
    return {k: float(v) for k, v in out.items()}


def _rad_fit(rg, y, ok, bw, plain=False):
    """plain=False: y = MSD(r)/MSD(far), plateau B bounded 0.5-1.5. plain=True: y = MSD(r) in A^2 (no far normalisation); the plateau B (A^2)
    is free between 0.5 and 3 x the mean of the outermost 3 fitted points. Radii (r50/r90/r95) are relative to the plateau, so they stay comparable."""
    from scipy.optimize import curve_fit
    r, v = rg[ok], y[ok]
    out = {}
    if len(r) < 6: return out
    tail = float(v[-3:].mean())
    A0 = float(np.clip(1 - v[:3].mean() / tail, 0.05, 1.0))
    n_ind = max((r[-1] - r[0]) / (2 * bw), 6.0)                         # kernel points overlap: independent points ~ span / (2 bw)
    for name in RAD_MODELS:
        sp = RAD_FIT[name]
        lo, hi, p0 = list(sp['lo']), list(sp['hi']), list(sp['p0'](A0))
        if plain:
            lo[-1], hi[-1], p0[-1] = 0.5 * tail, 3.0 * tail, tail          # B is the last parameter of every model
        try: popt, _ = curve_fit(sp['f'], r, v, p0=p0, bounds=(lo, hi), maxfev=5000)
        except Exception: continue
        mse = float(np.mean((sp['f'](r, *popt) - v) ** 2))
        par = dict(zip(sp['names'], map(float, popt)))
        out[name] = dict(par=par, rmse=float(np.sqrt(mse)), aic=float(n_ind * np.log(mse + 1e-12) + 2 * len(popt)), **_rad_radii(name, par))
    return out


def _rad_free(rg, y, ok, band):
    """Model-free numbers. r_noise = last grid radius where the profile is below 1 - band (band = 2 sd of the control profiles; 0 = nothing
    detectable). ell = integral of the deficit (1 - y) dr (A); V = integral of 4 pi r^2 (1 - y) dr (A^3) = volume of a fully frozen sphere that
    removes the same mobility; r_eff = its radius. The first observed value is held constant down to r = 0 (no Li there)."""
    r, v = rg[ok], y[ok]
    if len(r) < 3: return {}
    below = v < 1 - band[ok]
    d = 1 - v
    re, de = np.concatenate([[0.0], r]), np.concatenate([[d[0]], d])
    V = float(_trapz(4 * np.pi * re ** 2 * de, re))
    return dict(r_noise=float(r[below][-1]) if below.any() else 0.0, ell_A=float(_trapz(de, re)), V_A3=V,
                r_eff=float(np.sign(V) * (3 * abs(V) / (4 * np.pi)) ** (1 / 3)), r_max_used=float(r[-1]))


def _rad_pool(entries, rng=None, with_ctrl=True):
    """Sum the kernel sums of the folders of one dopant element; rng: resample the time blocks of every folder (block bootstrap)."""
    S = far = ctrl = 0
    nc = min(e['ctrl'].shape[0] for e in entries) if with_ctrl else 0
    for e in entries:
        nb = e['S'].shape[0]
        ix = np.arange(nb) if rng is None else rng.integers(0, nb, nb)
        S = S + e['S'][ix].sum(0); far = far + e['far'][ix].sum(0)
        if with_ctrl and nc: ctrl = ctrl + e['ctrl'][:nc][:, ix].sum(1)
    return S, far, (ctrl if (with_ctrl and nc) else None)


def _rad_eval(S, far, ctrl, rg, bw, min_neff, band=None, rmin=0.0, plain=False):
    """plain=True: y, yc and band are the plain MSD in A^2 (no division by the far MSD). The model-free numbers (r_noise, ell, V, r_eff) need a
    dimensionless deficit 1 - y, so they use y / scale with scale = mean of the outermost 5 valid points of the profile (its own plateau);
    r_noise / r_eff are therefore relative to the profile's outer level, not to the far-Li MSD."""
    import warnings
    with np.errstate(all='ignore'), warnings.catch_warnings():
        warnings.simplefilter('ignore')
        Mfar = far[1] / far[0] if far[0] > 0 else np.nan
        M = 1.0 if plain else Mfar
        y = S[1] / S[0] / M
        neff = S[0] ** 2 / S[2]
        ok = np.isfinite(y) & (neff >= min_neff)
        yc = sd = None
        if ctrl is not None:
            yc = ctrl[:, 1] / ctrl[:, 0] / M
            yc = np.where(ctrl[:, 0] ** 2 / ctrl[:, 2] >= min_neff, yc, np.nan)
            sd = np.nanstd(yc, axis=0, ddof=1) if yc.shape[0] >= 3 else np.full(len(rg), np.nan)
        if band is None:
            if plain:
                fb = 0.05 * float(np.nanmedian(y[ok])) if ok.any() else 0.05
                band = np.full(len(rg), fb) if sd is None else np.where(np.isfinite(sd), np.maximum(2 * sd, 0.02 * fb / 0.05), fb)
            else:
                band = np.full(len(rg), 0.05) if sd is None else np.where(np.isfinite(sd), np.maximum(2 * sd, 0.02), 0.05)
    if ok.sum() < 6: return None
    if plain:
        scale = float(np.nanmean(y[ok][-5:]))
        yn, bn, ycn = y / scale, band / scale, (None if yc is None else yc / scale)
    else:
        scale, yn, bn, ycn = 1.0, y, band, yc
    free = _rad_free(rg, yn, ok, bn)
    if ycn is not None and ycn.shape[0] >= 3 and free:
        nv = [f_['V_A3'] for f_ in (_rad_free(rg, row, np.isfinite(row), bn) for row in ycn) if f_]
        if len(nv) >= 3: free['r_eff_null'] = float((3 * 2 * np.std(nv, ddof=1) / (4 * np.pi)) ** (1 / 3))    # |r_eff| below this: not distinguishable from 0
    return dict(y=y, neff=neff, ok=ok, yc=yc, band=band, fits=_rad_fit(rg, y, ok & (rg >= rmin), bw, plain), free=free,
                plain=plain, M=float(Mfar), scale=scale)


def _rad_case(args):
    tag, entries, rg, bw, min_neff, nboot, seed, rmin = args[:8]
    plain = bool(args[8]) if len(args) > 8 else False
    try:
        base = _rad_eval(*_rad_pool(entries), rg, bw, min_neff, rmin=rmin, plain=plain)
        if base is None: return tag, None
        out = dict(base)
        out['ci'], out['ylo'], out['yhi'] = {}, np.full(len(rg), np.nan), np.full(len(rg), np.nan)
        if nboot > 0:
            rng = np.random.default_rng(seed)
            ys, vals = [], {}
            for _ in range(nboot):
                S, far, _c = _rad_pool(entries, rng, with_ctrl=False)
                ev = _rad_eval(S, far, None, rg, bw, min_neff, band=base['band'], rmin=rmin, plain=plain)
                if ev is None: continue
                ys.append(np.where(ev['ok'], ev['y'], np.nan))
                for m, ft in ev['fits'].items():
                    for k in ('A', 'B', 'r50', 'r90', 'r95'): vals.setdefault((m, k), []).append(ft['par'][k] if k in ('A', 'B') else ft[k])
                for k, v in ev['free'].items(): vals.setdefault(('free', k), []).append(v)
            import warnings
            with np.errstate(all='ignore'), warnings.catch_warnings():
                warnings.simplefilter('ignore')
                if len(ys) >= 10:
                    Y = np.array(ys)
                    out['ylo'], out['yhi'] = np.nanpercentile(Y, 16, axis=0), np.nanpercentile(Y, 84, axis=0)
                for k, v in vals.items():
                    if len(v) >= 10: out['ci'][k] = (float(np.percentile(v, 16)), float(np.percentile(v, 84)))
        return tag, out
    except Exception as e:
        return tag, f'{type(e).__name__}: {e}'


def _rad_best(r):
    return min(r['fits'], key=lambda m: r['fits'][m]['aic']) if r['fits'] else None


def _rad_csv_rows(tag, r, tau_ps, nf, ndop, meta):
    level, name, T, _ = tag
    nan = float('nan')
    base = dict(level=level, name=name, T=T, tau_ps=tau_ps, n_folders=nf, n_dopants=ndop, dopant_dopant_min_A=meta['dd'], half_box_A=meta['hb'],
                r_p99_A=meta['p99'])
    best, rows = _rad_best(r), []
    for m, ft in r['fits'].items():
        row = dict(base, model=m, best=int(m == best), rmse=ft['rmse'], aic=ft['aic'], r50=ft['r50'], r90=ft['r90'], r95=ft['r95'])
        for k in ('A', 'lam', 'r0', 'w', 'B'): row[k] = ft['par'].get(k, nan)
        for k in ('A', 'B', 'r50', 'r90', 'r95'):
            row[k + '_lo'], row[k + '_hi'] = r['ci'].get((m, k), (nan, nan))
        rows.append(row)
    if r['free']:
        row = dict(base, model='free', best='', **r['free'])
        for k in ('r_noise', 'ell_A', 'V_A3', 'r_eff'):
            row[k + '_lo'], row[k + '_hi'] = r['ci'].get(('free', k), (nan, nan))
        rows.append(row)
    return rows


def _rad_fmt(v, d=1, cap=None):
    """number; '*' = beyond the largest distance with enough Li (cap): extrapolated by the fit, not measured"""
    if v is None or not np.isfinite(v): return 'n/a'
    return f'{v:.{d}f}' + ('*' if cap is not None and v > cap else '')


def _rad_profile_fig(res, meta, rg, rfar, T, taus, labels, path, plain=False):
    """Rows = tau, columns = dopant element. Black = MSD(r)/MSD(far) (shaded: 16-84 % block bootstrap), grey band = control null (1 +- 2 sd),
    dashed = fits (bold = lowest AIC), dotted vertical = r95 of the best fit, grey vertical = far edge."""
    from matplotlib.lines import Line2D
    fig, axs = plt.subplots(len(taus), len(labels), figsize=(3.3 * len(labels), 2.9 * len(taus)), dpi=150, squeeze=False)
    for ri, tau in enumerate(taus):
        for ci, lab in enumerate(labels):
            ax = axs[ri, ci]
            ax.tick_params(labelsize=7)
            r = res.get((('element', lab, T, tau)))
            if ri == 0: ax.set_title(lab, fontsize=9)
            if ci == 0: ax.set_ylabel(f'tau = {tau:g} ps\n' + ('MSD(r)  (A^2)' if plain else 'MSD(r) / MSD(far)'), fontsize=8)
            if ri == len(taus) - 1: ax.set_xlabel('distance to the nearest dopant (A)', fontsize=8)
            if not isinstance(r, dict):
                ax.text(0.5, 0.5, 'n/a' if r is None else 'error', transform=ax.transAxes, ha='center', fontsize=8); continue
            ok = r['ok']
            ax.axhline(r['M'] if plain else 1, color='0.5', lw=0.6)                 # plain: dashed-grey level = MSD of the far Li (reference only, not used in the fit)
            ax.axvline(rfar, color='0.7', lw=0.6, ls=':')
            bo = np.isfinite(r['band']) & ok
            if plain and r['yc'] is not None and np.isfinite(r['yc']).any():
                with np.errstate(all='ignore'): cen = np.nanmean(r['yc'], axis=0)       # control profile (A^2) is the null centre
            else:
                cen = np.full(len(rg), r['M'] if plain else 1.0)
            bo &= np.isfinite(cen)
            ax.fill_between(rg[bo], (cen - r['band'])[bo], (cen + r['band'])[bo], color='0.88', zorder=0)
            both = ok & np.isfinite(r['ylo']) & np.isfinite(r['yhi'])
            if both.any(): ax.fill_between(rg[both], r['ylo'][both], r['yhi'][both], color='k', alpha=0.18, lw=0)
            ax.plot(rg[ok], r['y'][ok], 'k-', lw=1.1)
            best = _rad_best(r)
            for m, ft in r['fits'].items():
                sp = RAD_FIT[m]
                ax.plot(rg[ok], sp['f'](rg[ok], *[ft['par'][n] for n in sp['names']]), color=RAD_COL[m], ls='--' if m != best else '-',
                        lw=0.8 if m != best else 1.6)
            fr = r['free']
            cap = fr.get('r_max_used')
            if best and r['fits'][best]['r95'] > 0:
                ax.axvline(r['fits'][best]['r95'], color=RAD_COL[best], ls=':', lw=1)
            txt = [f'best: {best}', 'r95 ' + ' | '.join(f"{m[:3]} {_rad_fmt(r['fits'][m]['r95'], cap=cap)}" for m in r['fits']),
                   f"r_eff {_rad_fmt(fr.get('r_eff'))} (null {_rad_fmt(fr.get('r_eff_null'))})  r_noise {_rad_fmt(fr.get('r_noise'))}"]
            if best: txt.append(f"A {_rad_fmt(r['fits'][best]['par']['A'], 2)}  B {_rad_fmt(r['fits'][best]['par']['B'], 2)}")
            m_ = meta.get(('element', lab, T, tau))
            if m_: txt.append(f"folders {m_['nf']}, dopants {m_['nd']}")
            ax.text(0.98, 0.04, '\n'.join(txt), transform=ax.transAxes, ha='right', va='bottom', fontsize=5.5, family='monospace')
    h = [Line2D([], [], color='k', label=('MSD(r) (A^2), 16-84 % bootstrap' if plain else 'MSD(r) / MSD(far), 16-84 % bootstrap')),
         Line2D([], [], color='0.8', lw=6, label=('control null (mean +- 2 sd)' if plain else 'control null (1 +- 2 sd)'))]
    h += [Line2D([], [], color=RAD_COL[m], ls='--', label=f'fit {m}') for m in RAD_MODELS]
    h += [Line2D([], [], color='0.7', ls=':', label=f'far edge {rfar:g} A')]
    fig.legend(handles=h, loc='lower center', ncol=len(h), fontsize=6, frameon=False)
    fig.suptitle(f'T = {T} K   Li MSD vs distance to the nearest dopant, pooled over the folders of one dopant element'
                 + ('   [PLAIN MSD, no far normalisation; grey line = far-Li MSD]' if plain else ''), fontsize=8)
    fig.tight_layout(rect=(0, 0.04, 1, 0.97))
    fig.savefig(path); fig.savefig(os.path.splitext(path)[0] + '.svg')
    plt.close(fig)


RAD_MK = 'osv^D<>pP*'


def _rad_valid(r, m='sigmoid'):
    """The fit of model m of result r, only if it is a real measurement: r95 > 0, r95 inside the measured range (no '*'), depth A >= RAD_A_MIN
    (below that the radius is noise). Returns {r95, A, B, ci: {k: (lo, hi)}} or None."""
    if not isinstance(r, dict) or m not in r['fits']: return None
    ft = r['fits'][m]
    cap = r['free'].get('r_max_used', np.inf)
    if not (np.isfinite(ft['r95']) and 0 < ft['r95'] <= cap and ft['par']['A'] >= RAD_A_MIN): return None
    return dict(r95=ft['r95'], A=ft['par']['A'], B=ft['par']['B'], ci={k: r['ci'].get((m, k), (np.nan, np.nan)) for k in ('r95', 'A', 'B')})


def _rad_allT_fig(res, labels, Ts, taus, path, plain=False, model='sigmoid'):
    """Dopant comparison over ALL temperatures (and, with several taus, all lags) in one figure: columns = r95, depth A, plateau B;
    x = dopant element; colour = T; marker = tau (each tau is its own group inside the element slot). Only valid fits (see _rad_valid):
    no extrapolated radii, no A < RAD_A_MIN. Bars = 16-84 % block bootstrap."""
    from matplotlib.lines import Line2D
    multi = len(taus) > 1
    cols = {T: plt.cm.plasma(k / max(len(Ts) - 1, 1) * 0.85) for k, T in enumerate(Ts)}
    mk = {t: RAD_MK[i % len(RAD_MK)] for i, t in enumerate(taus)}
    nt, nT = len(taus), len(Ts)
    gw = 0.84 / nt
    fig, axs = plt.subplots(1, 3, figsize=(15 if multi else 11.5, 4.4), dpi=150)
    spec = (('r95', 'sigmoid r95 (A)'), ('A', 'depth of suppression A'),
            ('B', 'plateau B (A^2)' if plain else 'plateau B (x MSD(far))'))
    nval = ntot = 0
    for ai, (k, yl) in enumerate(spec):
        ax = axs[ai]
        for li, lab in enumerate(labels):
            if li: ax.axvline(li - 0.5, color='0.9', lw=0.6, zorder=0)
            for ti, tau in enumerate(taus):
                for Ti, T in enumerate(Ts):
                    r = res.get(('element', lab, T, tau))
                    if ai == 0 and isinstance(r, dict): ntot += 1
                    v = _rad_valid(r, model)
                    if not v: continue
                    if ai == 0: nval += 1
                    x = li - 0.42 + gw * (ti + 0.5) + (Ti - (nT - 1) / 2) * (gw * 0.7 / max(nT, 1))
                    lo, hi = v['ci'][k]
                    yerr = None if not (np.isfinite(lo) and np.isfinite(hi)) else [[max(v[k] - lo, 0)], [max(hi - v[k], 0)]]
                    ax.errorbar([x], [v[k]], yerr=yerr, marker=mk[tau], ms=5, ls='', capsize=1.5, lw=0.8, color=cols[T], mec='k', mew=0.4)
        ax.set_xticks(range(len(labels))); ax.set_xticklabels(labels, fontsize=8)
        ax.set_xlim(-0.5, len(labels) - 0.5)
        ax.set_ylabel(yl, fontsize=8); ax.tick_params(labelsize=7)
        if k == 'B' and plain and multi: ax.set_yscale('log')
        else: ax.set_ylim(bottom=0)
    h = [Line2D([], [], marker='o', ls='', color=cols[T], mec='k', mew=0.4, label=f'{T} K') for T in Ts]
    if multi: h += [Line2D([], [], marker=mk[t], ls='', color='0.6', mec='k', mew=0.4, label=f'tau = {t:g} ps') for t in taus]
    fig.legend(handles=h, loc='lower center', ncol=min(len(h), 12), fontsize=7, frameon=False)
    what = ('all temperatures, all lags' if multi else f'all temperatures, tau = {taus[0]:g} ps')
    fig.suptitle(f'{"plain MSD" if plain else "MSD / MSD(far)"} fit, sigmoid, {what}: valid fits only (r95 inside the measured range, A >= {RAD_A_MIN:g}); '
                 f'{nval} of {ntot} shown' + ('   [within an element: groups = tau, colour = T]' if multi else ''), fontsize=8)
    fig.tight_layout(rect=(0, 0.07, 1, 0.95))
    fig.savefig(path); fig.savefig(os.path.splitext(path)[0] + '.svg')
    plt.close(fig)


def _rad_summary_fig(res, labels, Ts, Tref, tau, model, path, plain=False):
    """T = Tref: (a) r95 per model, r_eff, r_noise per element; (b) A; (c) B; (d) r95 of --rad-model and r_eff vs T."""
    fig, axs = plt.subplots(2, 2, figsize=(10, 7.4), dpi=150)
    x = np.arange(len(labels))
    off = {'sigmoid': 0.0}

    def get(lab, T, m, k):
        r = res.get(('element', lab, T, tau))
        if not isinstance(r, dict): return np.nan, np.nan, np.nan
        if m == 'free': v = r['free'].get(k, np.nan)
        elif m in r['fits']: v = r['fits'][m]['par'][k] if k in ('A', 'B') else r['fits'][m][k]
        else: return np.nan, np.nan, np.nan
        lo, hi = r['ci'].get((m, k), (np.nan, np.nan))
        return v, lo, hi

    def bars(ax, m, k, dx, **kw):
        v = np.array([get(l, Tref, m, k) for l in labels]).T
        ax.errorbar(x + dx, v[0], yerr=[np.clip(v[0] - v[1], 0, None), np.clip(v[2] - v[0], 0, None)], ls='', capsize=2, **kw)

    ax = axs[0, 0]
    for m in RAD_SUM_MODELS: bars(ax, m, 'r95', off[m], marker='o', ms=4, color=RAD_COL[m], label=f'r95 {m}')
    bars(ax, 'free', 'r_eff', 0.4, marker='D', ms=4, color='k', label='r_eff (model-free)')
    bars(ax, 'free', 'r_noise', -0.4, marker='x', ms=5, color='0.4', label='r_noise (model-free)')
    ax.set_ylabel('radius (A)', fontsize=8); ax.set_title(f'radius per element, T = {Tref} K, tau = {tau:g} ps', fontsize=8); ax.legend(fontsize=6, frameon=False)
    ax = axs[0, 1]
    for m in RAD_SUM_MODELS: bars(ax, m, 'A', off[m], marker='o', ms=4, color=RAD_COL[m], label=m)
    ax.set_ylabel('A = max. relative MSD deficit at the dopant', fontsize=8); ax.set_title('depth of the suppression' + (' (plain MSD)' if plain else ''), fontsize=8); ax.legend(fontsize=6, frameon=False)
    ax = axs[1, 0]
    for m in RAD_SUM_MODELS: bars(ax, m, 'B', off[m], marker='o', ms=4, color=RAD_COL[m], label=m)
    if plain:
        ax.set_ylabel('B = fitted plateau MSD (A^2)', fontsize=8); ax.set_title('plateau of the fit (free, 0.5-3 x outer-point mean)', fontsize=8)
    else:
        ax.axhline(1, color='0.5', lw=0.6)
        ax.set_ylabel('B = plateau / MSD(far)', fontsize=8); ax.set_title('B > 1: MSD still rising beyond the far edge (radius > far edge)', fontsize=8)
    ax.legend(fontsize=6, frameon=False)
    for ax in (axs[0, 0], axs[0, 1], axs[1, 0]):
        ax.set_xticks(x); ax.set_xticklabels(labels, fontsize=8); ax.tick_params(labelsize=7)
    ax = axs[1, 1]                                  # temperature dependence: r95 of --rad-model only, valid fits only, one line per element
    n, span = len(labels), float(max(Ts) - min(Ts)) or 1.0
    drawn = False
    for i, lab in enumerate(labels):
        pts = [(T, _rad_valid(res.get(('element', lab, T, tau)), model)) for T in Ts]
        pts = [(T, v) for T, v in pts if v]
        if not pts: continue
        x = np.array([T for T, _ in pts], float) + (i - (n - 1) / 2) * 0.012 * span          # small offset so the bars do not overlap
        y = np.array([v['r95'] for _, v in pts])
        lo = np.array([v['ci']['r95'][0] for _, v in pts]); hi = np.array([v['ci']['r95'][1] for _, v in pts])
        yerr = [np.where(np.isfinite(lo), np.clip(y - lo, 0, None), 0), np.where(np.isfinite(hi), np.clip(hi - y, 0, None), 0)]
        ax.errorbar(x, y, yerr=yerr, marker='o', ms=5, lw=1.3, capsize=2, color=plt.cm.tab10(i % 10), label=lab)
        drawn = True
    if not drawn: ax.text(0.5, 0.5, 'no valid fits', transform=ax.transAxes, ha='center', fontsize=8)
    ax.set_xticks(Ts); ax.set_xlim(min(Ts) - 0.06 * span, max(Ts) + 0.06 * span); ax.set_ylim(bottom=0)
    ax.set_xlabel('T (K)', fontsize=8); ax.set_ylabel(f'{model} r95 (A)', fontsize=8)
    ax.set_title(f'temperature dependence, tau = {tau:g} ps (valid fits only)', fontsize=8)
    ax.tick_params(labelsize=7)
    if drawn: ax.legend(fontsize=7, frameon=False, ncol=min(n, 4))
    fig.tight_layout()
    fig.savefig(path); fig.savefig(os.path.splitext(path)[0] + '.svg')
    plt.close(fig)


def _rad_conc_fig(res, meta, flab, T, tau, path, plain=False):
    """Folder level (no bootstrap): sigmoid r95 (a) and depth A (b) vs number of dopant atoms in the folder, one line per dopant element.
    r95 is not drawn where A < RAD_A_MIN (the radius is noise there)."""
    fig, axs = plt.subplots(1, 2, figsize=(9, 3.6), dpi=150)
    pts = {}
    for tag, r in res.items():
        if tag[0] != 'folder' or tag[2] != T or tag[3] != tau or not isinstance(r, dict) or 'sigmoid' not in r['fits']: continue
        ft = r['fits']['sigmoid']
        pts.setdefault(flab.get(tag[1], '?'), []).append((meta[tag]['nd'], ft['r95'] if ft['par']['A'] >= RAD_A_MIN else np.nan, ft['par']['A'], tag[1]))
    for i, el in enumerate(sorted(pts, key=nat)):
        v = sorted(pts[el], key=lambda t: t[0])
        c = plt.cm.tab10(i % 10)
        x = [t[0] for t in v]
        axs[0].plot(x, [t[1] for t in v], '-o', ms=4, color=c, label=el)
        axs[1].plot(x, [t[2] for t in v], '-o', ms=4, color=c, label=el)
    axs[0].set_ylabel('sigmoid r95 (A)', fontsize=8); axs[0].set_title(f'radius vs doping, T = {T} K, tau = {tau:g} ps (A < {RAD_A_MIN:g} not drawn)', fontsize=8)
    axs[1].set_ylabel('sigmoid depth A', fontsize=8); axs[1].set_title('depth vs doping' + (' (plain MSD)' if plain else ''), fontsize=8)
    for ax in axs:
        ax.set_xlabel('dopant atoms in the folder', fontsize=8); ax.tick_params(labelsize=7); ax.legend(fontsize=6, frameon=False)
        ax.xaxis.set_major_locator(plt.MaxNLocator(integer=True))
    fig.tight_layout()
    fig.savefig(path); fig.savefig(os.path.splitext(path)[0] + '.svg')
    plt.close(fig)


def make_radius(rad, outdir, a, p, plain=False):
    """Pool the kernel sums per dopant element (and per folder), fit, bootstrap, write csv + figures into outdir (NOT added to the report).
    plain=True: the same kernel sums, but the fitted profile is the plain MSD(r) in A^2 (not divided by MSD(far)); see _rad_eval."""
    rfar = float(a.rad_far) if a.rad_far else float(p['edges'][-1])
    ents, fents, flab, seen, rg = {}, {}, {}, set(), None
    for (gi, f), results in rad.items():
        if os.path.abspath(f) in seen: continue
        seen.add(os.path.abspath(f))
        for r in results:
            g_ = np.array(r['rg'], float)
            if rg is None: rg = g_
            if len(g_) != len(rg) or not np.allclose(g_, rg): warn(f'radius: grid of {f} differs, skipped'); continue
            for lg in r['lags']:
                S = np.array(lg['S'], float)
                e = dict(S=S, far=np.array(lg['far'], float), ctrl=np.array(lg['ctrl'], float).reshape((-1,) + S.shape), tau_ps=lg['tau_ps'], meta=r['meta'])
                ents.setdefault((r['label'], r['T'], lg['tau_req']), []).append(e)
                fents[(r['folder'], r['T'], lg['tau_req'])] = [e]
                flab[r['folder']] = r['label']
    if not ents: return None
    os.makedirs(outdir, exist_ok=True)
    tasks, info = [], {}
    for i, (k, lst) in enumerate(sorted(ents.items(), key=lambda kv: (nat(kv[0][0]), kv[0][1], kv[0][2]))):
        tag = ('element',) + k
        tasks.append((tag, lst, rg, a.rad_bw, a.rad_min_neff, a.rad_boot, a.ctrl_seed * 1000 + i, a.rad_rmin, plain))
        info[tag] = lst
    for i, (k, lst) in enumerate(sorted(fents.items(), key=lambda kv: (nat(kv[0][0]), kv[0][1], kv[0][2]))):
        tag = ('folder',) + k
        tasks.append((tag, lst, rg, a.rad_bw, a.rad_min_neff, 0, 0, a.rad_rmin, plain))
        info[tag] = lst
    print(f'[radius{" plain MSD" if plain else ""}] fitting {len(tasks)} cases (bootstrap {a.rad_boot} on the element level)')
    if a.jobs > 1 and len(tasks) > 1:
        from concurrent.futures import ProcessPoolExecutor
        with ProcessPoolExecutor(max_workers=a.jobs) as ex: out = list(ex.map(_rad_case, tasks))
    else:
        out = [_rad_case(t) for t in tasks]
    res, meta, rows, prof = {}, {}, [], []
    for tag, r in out:
        lst = info[tag]
        if not isinstance(r, dict):
            warn(f'radius {tag}: ' + ('too few Li-origin samples near the dopant (see --rad-min-neff)' if r is None else r)); res[tag] = r; continue
        res[tag] = r
        dd = [e['meta']['dd_min'] for e in lst if np.isfinite(e['meta']['dd_min'])]
        mt = dict(dd=float(np.median(dd)) if dd else float('nan'), hb=float(min(e['meta']['half_box'] for e in lst)),
                  p99=float(np.median([e['meta']['r_p99'] for e in lst])))
        nd = int(sum(e['meta']['n_dop'] for e in lst))
        meta[tag] = dict(nf=len(lst), nd=nd)
        tau_ps = float(np.mean([e['tau_ps'] for e in lst]))
        rows += _rad_csv_rows(tag, r, tau_ps, len(lst), nd, mt)
        fitc = {m: RAD_FIT[m]['f'](rg, *[r['fits'][m]['par'][n] for n in RAD_FIT[m]['names']]) for m in r['fits']}
        for j in np.where(r['ok'])[0]:
            yk = 'msd_A2' if plain else 'ratio'
            row = dict(level=tag[0], name=tag[1], T=tag[2], tau_ps=tau_ps, r_A=float(rg[j]), n_eff=float(r['neff'][j]), null_band=float(r['band'][j]),
                       far_msd_A2=float(r['M']))
            row[yk], row[yk + '_lo'], row[yk + '_hi'] = float(r['y'][j]), float(r['ylo'][j]), float(r['yhi'][j])
            if r['yc'] is not None: row['ctrl_mean'] = float(np.nanmean(r['yc'][:, j])) if np.isfinite(r['yc'][:, j]).any() else float('nan')
            for m in fitc: row['fit_' + m] = float(fitc[m][j])
            prof.append(row)
    if not rows: return None
    with open(os.path.join(outdir, 'radius_summary.csv'), 'w', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=RAD_COLS, restval='nan'); w.writeheader(); w.writerows(rows)
    yk = 'msd_A2' if plain else 'ratio'
    pc = ['level', 'name', 'T', 'tau_ps', 'r_A', yk, yk + '_lo', yk + '_hi', 'far_msd_A2', 'n_eff', 'null_band', 'ctrl_mean'] + ['fit_' + m for m in RAD_MODELS]
    with open(os.path.join(outdir, 'radius_profiles.csv'), 'w', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=pc, restval='nan'); w.writeheader(); w.writerows(prof)
    labels = sorted({t[1] for t in res if t[0] == 'element'}, key=nat)
    Ts = sorted({t[2] for t in res if t[0] == 'element'})
    taus = sorted({t[3] for t in res if t[0] == 'element'})
    tau_of = {}
    figs = []
    for T in Ts:
        pth = os.path.join(outdir, f'radius_profiles_T{T}K.png')
        _rad_profile_fig(res, meta, rg, rfar, T, taus, labels, pth, plain); figs.append(pth)
    Tref = a.rad_tref if a.rad_tref in Ts else max(Ts)
    for tau in taus:
        pth = os.path.join(outdir, f'radius_summary_tau{tau:g}ps.png')
        _rad_summary_fig(res, labels, Ts, Tref, tau, a.rad_model, pth, plain); figs.append(pth)
        pth = os.path.join(outdir, f'radius_vs_conc_T{Tref}K_tau{tau:g}ps.png')
        _rad_conc_fig(res, meta, flab, Tref, tau, pth, plain); figs.append(pth)
        pth = os.path.join(outdir, f'radius_allT_tau{tau:g}ps.png')                      # all temperatures, one tau
        _rad_allT_fig(res, labels, Ts, [tau], pth, plain, a.rad_model); figs.append(pth)
        for lab in labels:
            r = res.get(('element', lab, Tref, tau))
            if not isinstance(r, dict): print(f'[radius] T={Tref} tau={tau:g} {lab}: n/a'); continue
            b = _rad_best(r)
            cap = r['free'].get('r_max_used')
            print(f'[radius{" plain" if plain else ""}] T={Tref} tau={tau:g} {lab:>3}: best={b}  r95 ' + ' / '.join(f"{m[:3]} {_rad_fmt(r['fits'][m]['r95'], cap=cap)}" for m in r['fits'])
                  + f"   r_eff {_rad_fmt(r['free'].get('r_eff'))} (null {_rad_fmt(r['free'].get('r_eff_null'))})  r_noise {_rad_fmt(r['free'].get('r_noise'))}"
                  + (f"   A {_rad_fmt(r['fits'][b]['par']['A'], 2)}  B {_rad_fmt(r['fits'][b]['par']['B'], 2)}" if b else ''))
    pth = os.path.join(outdir, 'radius_allT_alltau.png')                                # all temperatures AND all lags
    _rad_allT_fig(res, labels, Ts, taus, pth, plain, a.rad_model); figs.append(pth)
    print('[radius] * = beyond the measured range (extrapolated by the fit); r_eff below its null value is not distinguishable from 0')
    print(f'[radius] {outdir}/radius_summary.csv radius_profiles.csv ' + ' '.join(os.path.basename(f) for f in figs))
    return dict(figs=figs)


def dev_blocks(entries, edges):
    """entries [(folder label, records)] of one group -> tables after the comparison."""
    names = lio.shell_names(np.array(edges, float))
    systems = []
    for lab, recs in entries:
        d = {(r['T'], r['species'], r['ref'], r['shell']): r for r in recs}
        refs = ['any'] + [r['ref'] for r in recs if r['species'] == 'dop' and r['ref'] != 'any' and r['T'] == recs[0]['T']]
        systems.append(dict(name=lab, d=d, refs=refs, Ts={r['T'] for r in recs}))
    Ts = sorted(set().union(*[s['Ts'] for s in systems]))
    B = [('h', 1, 'Site deviation from pure LLZO', True)]
    B.append(('p', 'Deviation (Å) = distance of an atom from its matched site in the pure Ia-3d LLZO lattice (rigidly aligned on the '
                      'La/Zr framework, scaled to the MD box). La, Zr, O: time-mean position matched one-to-one to the closest free site; '
                      'dopant: closest site of the stated sublattice (time-mean position); Li: mean over the frames spent in the shell, nearest free Li site '
                      '(shell of a Li = its distance to the dopant in that frame, dopant position smoothed over --dop-smooth-ps, default 1 ps). '
                      'Cell = median ± half-width of the 95% bootstrap CI over atoms (n = atoms; for Li, atom-shell pairs). '
                      'Shell = distance to the dopant (average = nearest dopant; Ca1 ... = that dopant atom); "all" = every atom of the species.'))
    for sp, title in (('Li', 'Li'), ('La', 'La'), ('Zr', 'Zr'), ('O', 'O'), ('dop', 'Dopant atom')):
        if not any(k[1] == sp for s in systems for k in s['d']): continue
        B.append(('h', 2, f'{title}: deviation from pure-LLZO site', False))
        rows = ['all'] if sp == 'dop' else names + ['all']
        B.append(('t', grid_data(f'{title} deviation from pure-LLZO site (Å): median ± 95% CI half-width (n)',
                   [(s, s['name'], s['refs']) for s in systems], Ts, rows, lambda r: 'all' if r == 'all' else summ.shell_name(r),
                   lambda s, T, a, r, sp=sp: dev_cell(s['d'].get((T, sp, a, r))))))
    return B


# ---------------------------------------------------------------- LaTeX export: figures1/ + <out>.tex
FIGDIR = 'figures1'


class FigStore:
    """Copies every figure of the report into <outdir>/figures1 under a unique LaTeX-safe name.
    .svg is copied whenever it exists; a .pdf is made from it (rsvg-convert or inkscape, if installed) because pdflatex cannot
    read svg; the .png is copied only when no vector version could be made. \\includegraphics without extension takes the .pdf first."""
    def __init__(self, outdir):
        self.dir = os.path.join(outdir, FIGDIR)
        os.makedirs(self.dir, exist_ok=True)
        self.names, self.used, self.warned = {}, set(), False
        self.conv = shutil.which('rsvg-convert') or shutil.which('inkscape')

    def _pdf(self, svg, pdf):
        if not self.conv: return False
        if os.path.exists(pdf) and os.path.getmtime(pdf) >= os.path.getmtime(svg): return True
        if os.path.basename(self.conv).startswith('rsvg'): cmd = [self.conv, '-f', 'pdf', '-o', pdf, svg]
        else: cmd = [self.conv, svg, '--export-type=pdf', '--export-filename=' + pdf]
        r = subprocess.run(cmd, capture_output=True)
        return r.returncode == 0 and os.path.exists(pdf)

    def get(self, path, prefix=''):
        path = os.path.abspath(path)
        if path in self.names: return self.names[path]
        name = re.sub(r'[^A-Za-z0-9_-]+', '_', f'{prefix}_{os.path.splitext(os.path.basename(path))[0]}').strip('_')
        base, k = name, 1
        while name in self.used:
            k += 1; name = f'{base}_{k}'
        self.used.add(name)
        self.names[path] = name
        dst = os.path.join(self.dir, name)
        svg = os.path.splitext(path)[0] + '.svg'
        vector = False
        if os.path.exists(svg):
            shutil.copy2(svg, dst + '.svg')
            vector = self._pdf(dst + '.svg', dst + '.pdf')
        if not vector:
            shutil.copy2(path, dst + os.path.splitext(path)[1])
            if os.path.exists(svg) and not self.warned:
                self.warned = True
                warn('svg -> pdf conversion not possible (install rsvg-convert: sudo apt install librsvg2-bin, or inkscape); png used in the tex')
        return name


def tex_esc(s):
    return re.sub(r'([&%$#_{}])', r'\\\1', str(s))


def tex_blocks(fs, prefix, paths, row, max_h):
    """-> latex lines for one STEPS block ('row' = side by side, 'each' = one per line)."""
    if row:
        w = f'{0.97 / len(paths):.3f}'
        im = [f'\\includegraphics[width={w}\\linewidth,height={max_h}\\textheight,keepaspectratio]{{{fs.get(p, prefix)}}}' for p in paths]
        return ['\\noindent' + '\\hfill'.join(im) + '\\par', '']
    return [f'{{\\centering\\includegraphics[width=\\linewidth,height={max_h}\\textheight,keepaspectratio]{{{fs.get(p, prefix)}}}\\par}}'
            for p in paths] + ['']


def tex_sitedev(fs, prefix, name, rows, title):
    """one page, one row per T, every image the same height (same sizing as add_sitedev_page)."""
    nrow, ncol = len(rows), len(rows[0])
    asp = [[aspect(p) for p in r] for r in rows]
    ca = [max(a[i] for a in asp) for i in range(ncol)]
    hh = min((USABLE_W - 0.6 * ncol) / sum(ca), (PAGE_H - 2 * MARGIN - 2.0) / nrow - 0.35)
    out = ['\\clearpage', f'\\subsection*{{{tex_esc(title)}: {tex_esc(name)}}}']
    for paths in rows:
        out.append('{\\centering ' + '\\hspace{0.2cm}'.join(f'\\includegraphics[height={hh:.2f}cm]{{{fs.get(pt, prefix)}}}' for pt in paths) + '\\par}')
    return out + ['']


_TEX_ESC = {'\\': '\\textbackslash{}', '&': '\\&', '%': '\\%', '$': '\\$', '#': '\\#', '_': '\\_', '{': '\\{', '}': '\\}',
            '~': '\\textasciitilde{}', '^': '\\textasciicircum{}', '|': '\\textbar{}'}
_TEX_UNI = {'±': '$\\pm$', 'Å': '\\AA{}', '–': '--', '—': '---', 'τ': '$\\tau$', '≥': '$\\geq$', '≤': '$\\leq$', '→': '$\\rightarrow$',
            '×': '$\\times$', '²': '$^2$', '·': '$\\cdot$', '≈': '$\\approx$', 'Δ': '$\\Delta$', 'σ': '$\\sigma$', 'μ': '$\\mu$',
            'µ': '$\\mu$', '−': '$-$', '…': '\\ldots{}'}


def tex_txt(s):
    s = ''.join(_TEX_ESC.get(c, c) for c in str(s))
    for k, v in _TEX_UNI.items(): s = s.replace(k, v)
    if any(ord(c) > 127 for c in s):
        warn('tex: unmapped character(s) ' + ''.join(sorted({c for c in s if ord(c) > 127})))
        s = ''.join('?' if ord(c) > 127 else c for c in s)
    return s


def tex_grid(d):
    nd, ns = len(d['atoms']), len(d['rowlabels'])
    nc = 2 + nd
    pt = 8 if nc <= 10 else 7 if nc <= 16 else 6
    wd = (USABLE_W - 3.6) / nd - 0.17
    col = lambda w: f'>{{\\centering\\arraybackslash}}p{{{w:.2f}cm}}|'
    L = ['\\begingroup', f'\\fontsize{{{pt}}}{{{pt + 1.5}}}\\selectfont\\setlength{{\\tabcolsep}}{{2pt}}',
         '\\begin{longtable}{|' + col(1.55) + col(1.75) + col(wd) * nd + '}', '\\hline',
         '\\multirow{3}{1.5cm}{\\centering\\textbf{Temperature (K)}} & \\multirow{3}{1.7cm}{\\centering\\textbf{Shell}} & '
         f'\\multicolumn{{{nd}}}{{c|}}{{\\parbox{{{wd * nd:.2f}cm}}{{\\centering\\textbf{{{tex_txt(d["defn"])}}}}}}} \\\\',
         f'\\cline{{3-{nc}}}',
         '& & ' + ' & '.join(f'\\multicolumn{{{n}}}{{c|}}{{\\textbf{{{tex_txt(lab)}}}}}' for lab, n in d['syshead']) + ' \\\\',
         f'\\cline{{3-{nc}}}', '& & ' + ' & '.join(f'\\textbf{{{tex_txt(a)}}}' for a in d['atoms']) + ' \\\\', '\\hline', '\\endhead']
    for ti, T in enumerate(d['Ts']):
        for i, lab in enumerate(d['rowlabels']):
            first = f'\\multirow{{{ns}}}{{1.5cm}}{{\\centering {T}}}' if i == 0 else ''
            L.append(' & '.join([first, tex_txt(lab)] + [tex_txt(x) for x in d['cells'][ti][i]]) + ' \\\\')
            L.append(f'\\cline{{2-{nc}}}' if i < ns - 1 else '\\hline')
    return L + ['\\end{longtable}', '\\endgroup', '']


def tex_tables(blocks):
    L = []
    for b in blocks:
        if b[0] == 'h':
            if b[3]: L.append('\\clearpage')
            L += [('\\section*{%s}' if b[1] == 1 else '\\subsection*{%s}') % tex_txt(b[2]), '']
        elif b[0] == 'p': L += [tex_txt(b[1]), '']
        else: L += tex_grid(b[1])
    return L


def write_tex(groups, sd, rm, outpath, tabs=None, rho=None, pool=None):
    """<out>.tex next to the report + figures1/ with every figure. Text around the figures is yours to add. Tables = the docx tables."""
    outdir = os.path.dirname(os.path.abspath(outpath))
    base = os.path.splitext(os.path.basename(outpath))[0]
    fs = FigStore(outdir)
    rho = rho or {}
    L = ['\\documentclass[10pt]{article}', '\\usepackage[T1]{fontenc}', '\\usepackage[a4paper,landscape,margin=1.27cm]{geometry}',
         '\\usepackage{graphicx}', '\\usepackage{longtable,multirow,array}', f'\\graphicspath{{{{{FIGDIR}/}}}}', '\\setlength{\\parindent}{0pt}', '\\setlength{\\parskip}{4pt}',
         '', '\\begin{document}', '']
    first = True
    for gi, g in enumerate(groups):
        for f in g['folders']:
            if not os.path.isdir(f): continue
            pre = fname(f)
            if not first: L.append('\\clearpage')
            first = False
            L += [f'\\section*{{{tex_esc(pre)}}}', '']
            first_img = True
            for layout, sub, pat in STEPS:
                paths = find(f, sub, pat)
                if not paths:
                    L += [f'% missing: {os.path.join(f, sub, "**", pat)}', '']
                    continue
                L += tex_blocks(fs, pre, paths, layout == 'row', 0.68 if first_img else 0.9)
                if first_img and rm.get((gi, f)): L += tex_blocks(fs, pre, [rm[(gi, f)]], False, 0.9)     # RMSD after the very first image
                first_img = False
            if rho.get((gi, f)): L += tex_blocks(fs, pre, rho[(gi, f)], False, 0.9)       # local-density bins, after the msd05 image
            r = sd.get((gi, f))
            if r:
                L += tex_sitedev(fs, pre, pre, [[x['violin']] + x['projs'] for x in r[0]], 'Site deviation from pure LLZO')
                L += tex_sitedev(fs, pre, pre, [x['full'] for x in r[0]], 'Whole-cell projection')
        L += tex_tables((tabs or {}).get(gi, []))
    if pool:
        L += ['\\clearpage', '\\section*{Shell density vs MSD and $H_R$ (pooled over folders)}', '', tex_txt(pool['text']), '']
        for pth in pool['figs']: L += tex_blocks(fs, 'shellpool', [pth], False, 0.8)
    L.append('\\end{document}')
    tex = os.path.join(outdir, base + '.tex')
    with open(tex, 'w') as fh: fh.write('\n'.join(L) + '\n')
    return tex


def make_tables(groups, sd, p, cm):
    """{group index: blocks} comparison + deviation tables (shared by docx and tex)."""
    tabs = {}
    for gi, g in enumerate(groups):
        ok, dev, B = [], [], []
        for f in g['folders']:
            if not os.path.isdir(f): continue
            ok.append(f)
            r = sd.get((gi, f))
            if r: dev.append((fname(f), r[1]))
        if len(ok) < 2:
            warn(f'group {gi + 1}: <2 valid folders, comparison skipped')
        else:
            systems = load_systems(ok, gi, cm)
            if len(systems) < 2: warn(f'group {gi + 1}: <2 systems with readable 05/07 logs, comparison skipped')
            else: B += compare_blocks(systems)
        if dev: B += dev_blocks(dev, p['edges'])
        tabs[gi] = B
    return tabs


def build_report(groups, sd, p, rm=None, tabs=None, rho=None, pool=None):
    rm, tabs, rho = rm or {}, tabs or {}, rho or {}
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(PAGE_W), Cm(PAGE_H)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(MARGIN)

    first = True
    for gi, g in enumerate(groups):
        for f in g['folders']:
            if not os.path.isdir(f):
                warn(f'not a folder, skipped: {f}'); continue
            r = sd.get((gi, f))
            add_folder(doc, f, page_break=not first, sd=r[0] if r else None, rmsd=rm.get((gi, f)), rho=rho.get((gi, f)))
            first = False
        render_docx(doc, tabs.get(gi, []))
    if pool:
        h = doc.add_heading('Shell density vs MSD and H_R (pooled over folders)', 1)
        h.paragraph_format.page_break_before = True
        doc.add_paragraph(pool['text'])
        for pth in pool['figs']: add_images(doc, [pth], False, MAX_H - 3.0)
    return doc


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('items', nargs='+', help='groups: {A B C} {D E Ga=Li24d}')
    ap.add_argument('-o', '--out', default='report.docx')
    ap.add_argument('--scripts', help='folder with llzo_io.py, llzo_sites.py ... (default: ./with_flag next to this file)')
    ap.add_argument('--cif', default=CIF_DEFAULT, help='pristine Ia-3d garnet CIF (Li, La, Zr, O)')
    ap.add_argument('--T', type=int, nargs='*', help='temperatures for the site deviation (default: all dumps found)')
    ap.add_argument('--tmin', type=float, help='analysis window start, ps (default: tmin_ps of the settings)')
    ap.add_argument('--shells', type=float, nargs='*', default=list(lio.DEFAULT_SHELLS), help='shell edges, A (as 05)')
    ap.add_argument('--every', type=int, default=10, help='use every Nth frame of the window')
    ap.add_argument('--window', type=float, default=8.0, help='projection half-width in-plane, A')
    ap.add_argument('--slab', type=float, default=1.5, help='projection slab half-thickness, A')
    ap.add_argument('--jobs', type=int, default=1, help='parallel (folder, T) computations')
    ap.add_argument('--tau-ps', type=float, default=1.0, help='return window for committed hops, ps (frames = tau / dump frame spacing)')
    ap.add_argument('--rho-r', type=float, default=4.0, help='local Li density: neighbours within this radius, A')
    ap.add_argument('--rho-bins', type=int, default=5, help='local Li density: number of fixed-width bins of relative local density (0.1-1.5)')
    ap.add_argument('--rho-lags', type=float, nargs='*', default=[1.0, 5.0, 10.0], help='local Li density: tau values, ps')
    ap.add_argument('--pool-lags', type=float, nargs='*', default=[2.0, 5.0, 10.0], help='shell density pool: tau values, ps')
    ap.add_argument('--pool-min-n', type=float, default=3.0, help='shell density pool: shells with fewer Li (mean per origin) are pale / not correlated')
    ap.add_argument('--ctrl-n', type=int, default=10, help='shell density pool: number of non-dopant control reference points')
    ap.add_argument('--ctrl-ref', choices=['li', 'random'], default='li', help='control points: positions of random Li atoms at frame 0 (li) or random points')
    ap.add_argument('--ctrl-seed', type=int, default=0, help='shell density pool: random seed of the control points')
    ap.add_argument('--rad-lags', type=float, nargs='*', default=[2.0, 5.0, 10.0, 20.0, 30.0], help='radius of influence: tau values, ps (a tau needs a window of at least 2 x blocks x tau frames, else it is skipped with a warning)')
    ap.add_argument('--rad-bw', type=float, default=0.5, help='radius of influence: Gaussian kernel width, A')
    ap.add_argument('--rad-step', type=float, default=0.25, help='radius of influence: grid step, A')
    ap.add_argument('--rad-rmax', type=float, default=12.0, help='radius of influence: largest distance on the grid, A')
    ap.add_argument('--rad-far', type=float, help='radius of influence: far-Li normaliser = Li farther than this from every dopant (default: last shell edge)')
    ap.add_argument('--rad-boot', type=int, default=200, help='radius of influence: block-bootstrap replicates (element level)')
    ap.add_argument('--rad-min-neff', type=float, default=100.0, help='radius of influence: grid points with fewer kernel-weighted Li-origins are dropped')
    ap.add_argument('--rad-rmin', type=float, default=2.5, help='radius of influence: fit only r >= this, A (0 = whole profile)')
    ap.add_argument('--rad-tref', type=int, help='radius of influence: T of the summary figure and the printed table (default: highest T)')
    ap.add_argument('--rad-model', choices=['exp', 'gauss', 'sigmoid'], default='sigmoid', help='radius of influence: model of the r95 vs T panel')
    ap.add_argument('--radius-only', action='store_true', help='compute only the radius of influence (no docx / tex / pdf, nothing else recomputed)')
    ap.add_argument('--no-radius', action='store_true', help='skip the radius of influence')
    ap.add_argument('--no-plain', action='store_true', help='radius of influence: skip the extra plain-MSD fit (radius_plainmsd/)')
    ap.add_argument('--recompute', action='store_true', help='ignore cached sitedev, rmsd and local-density results')
    lio.add_config_args(ap)
    a = ap.parse_args()
    groups = parse_groups(a.items)
    if not groups: sys.exit('no groups given')
    for g in groups: g['folders'] = expand(g['folders'])

    p = dict(cif=a.cif, every=a.every, tmin=a.tmin, edges=list(a.shells), minframes=3, boot=1000, window=a.window, slab=a.slab)
    outdir_rad = os.path.join(os.path.dirname(os.path.abspath(a.out)), 'radius')
    outdir_rad_plain = os.path.join(os.path.dirname(os.path.abspath(a.out)), 'radius_plainmsd')

    def radius_both():
        # kernel sums are computed (or read from the cache) once; the ratio fit and the plain-MSD fit both start from them
        rad = run_radius(groups, a, p)
        make_radius(rad, outdir_rad, a, p)
        if not a.no_plain: make_radius(rad, outdir_rad_plain, a, p, plain=True)

    if a.radius_only:
        radius_both()
        return
    sd = {}
    if any(g['sites'] for g in groups):
        if not os.path.exists(a.cif): warn(f'CIF not found, site deviation skipped: {a.cif}')
        else: sd = run_sitedev(groups, a, p)
    rm = run_rmsd(groups, a, p)
    cm = run_commit(groups, a, p)
    rh = run_rho(groups, a, p)
    sp = run_shellpool(groups, a, p)
    pool = make_pool(sp, os.path.join(os.path.dirname(os.path.abspath(a.out)), 'shellpool'), a.pool_min_n)
    if not a.no_radius: radius_both()
    tabs = make_tables(groups, sd, p, cm)

    global SVG_MODE
    SVG_MODE = False
    doc = build_report(groups, sd, p, rm, tabs, rh, pool)
    doc.save(a.out)
    print(f'wrote {os.path.abspath(a.out)}  ({len(groups)} group(s), png)')
    tex = write_tex(groups, sd, rm, a.out, tabs, rh, pool)
    print(f'wrote {tex}  (figures in {os.path.join(os.path.dirname(tex), FIGDIR)}/)')
    # pdf: same report built again with svg pictures, converted by LibreOffice (vector)
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice: sys.exit('soffice/libreoffice not found -> no pdf')
    import tempfile
    SVG_MODE = True
    base = os.path.splitext(os.path.basename(a.out))[0]
    with tempfile.TemporaryDirectory() as tmp:
        tdoc = os.path.join(tmp, base + '.docx')
        build_report(groups, sd, p, rm, tabs, rh, pool).save(tdoc)
        subprocess.run([soffice, '--headless', '--convert-to', 'pdf', '--outdir', tmp, tdoc], check=True)
        pdf = os.path.join(os.path.dirname(os.path.abspath(a.out)), base + '.pdf')
        shutil.move(os.path.join(tmp, base + '.pdf'), pdf)
    if SVG_MISSING: print(f'note: {len(SVG_MISSING)} image(s) have no .svg next to the .png -> PNG used in pdf, e.g. {SVG_MISSING[0]}', file=sys.stderr)
    print(f'wrote {pdf}  (svg)')


if __name__ == '__main__':
    main()
